from __future__ import annotations

# PATCH MARKER: SECTION_COMPARISON_EXPORT_FUNCTION_PRESENT
# PATCH MARKER: REPORT_COLUMN_SIDE_BARS_SUPPORTED
# PATCH MARKER: INITIAL_STRUCTURE_AND_UTILIZATION_EXPORT_PRESENT
# PATCH MARKER: UTILIZATION_EVOLUTION_MAPS_PRESENT

import os
from dataclasses import asdict, is_dataclass
from datetime import date
from typing import Any, Optional

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt


# ============================================================
# Robust helpers
# ============================================================


def _first_col(df: Optional[pd.DataFrame], names: list[str]) -> Optional[str]:
    if df is None or getattr(df, "empty", True):
        return None
    lookup = {str(c).strip().lower(): c for c in df.columns}
    for name in names:
        key = str(name).strip().lower()
        if key in lookup:
            return lookup[key]
    return None


def _safe_float(value: Any, default: float = np.nan) -> float:
    try:
        if pd.isna(value):
            return default
        value = float(value)
        return value if np.isfinite(value) else default
    except Exception:
        return default


def _cost_to_dict(cost_obj: Any) -> dict[str, Any]:
    if cost_obj is None:
        return {}
    if is_dataclass(cost_obj):
        return asdict(cost_obj)
    if isinstance(cost_obj, dict):
        return dict(cost_obj)
    out = {}
    for key in ["concrete_m3", "rebar_kg", "formwork_m2", "material_cost", "labor_cost", "total_cost"]:
        if hasattr(cost_obj, key):
            out[key] = getattr(cost_obj, key)
    return out


def _max_from_cols(df: Optional[pd.DataFrame], cols: list[str]) -> float:
    if df is None or getattr(df, "empty", True):
        return np.nan
    for col in cols:
        real = _first_col(df, [col])
        if real is not None:
            vals = pd.to_numeric(df[real], errors="coerce").abs().dropna()
            if not vals.empty:
                return float(vals.max())
    return np.nan


# ============================================================
# Drift extraction: source of truth is df_static_drifts_before
# ============================================================


def drift_rows_from_any(obj: Any, default_case: str = "") -> pd.DataFrame:
    """Normalize drift data to Story, Case, Drift_Ratio rows.

    This deliberately removes envelope/non-story rows. It accepts dataframes,
    dictionaries, lists and common optimizer nested objects.
    """
    rows: list[dict[str, Any]] = []

    def add(story: Any, drift: Any, case: Any = "") -> None:
        st = pd.to_numeric(pd.Series([story]), errors="coerce").iloc[0]
        dr = pd.to_numeric(pd.Series([drift]), errors="coerce").iloc[0]
        if pd.notna(st) and pd.notna(dr):
            rows.append({
                "Story": int(st),
                "Case": str(case or default_case or ""),
                "Drift_Ratio": abs(float(dr)),
            })

    def walk(x: Any, case: str = "") -> None:
        if x is None:
            return
        if isinstance(x, pd.DataFrame):
            if x.empty:
                return
            story_col = _first_col(x, ["Story", "story", "Storey", "storey", "Level", "Floor"])
            drift_col = _first_col(x, [
                "Drift_Ratio", "drift_ratio", "Story_Drift_Ratio", "Max_Drift_Ratio",
                "Drift", "drift", "ratio", "Max drift ratio", "MaxDriftRatio",
            ])
            case_col = _first_col(x, ["Case", "case", "Combination", "Load Combination", "Record", "EQ_Case"])
            if story_col is not None and drift_col is not None:
                for _, rr in x.iterrows():
                    add(rr.get(story_col), rr.get(drift_col), rr.get(case_col, case) if case_col else case)
            return
        if isinstance(x, dict):
            story_key = next((k for k in ["Story", "story", "Storey", "storey", "Level", "Floor"] if k in x), None)
            drift_key = next((k for k in ["Drift_Ratio", "drift_ratio", "Story_Drift_Ratio", "Max_Drift_Ratio", "Drift", "drift", "ratio"] if k in x), None)
            case_key = next((k for k in ["Case", "case", "Combination", "Record", "EQ_Case"] if k in x), None)
            if story_key is not None and drift_key is not None:
                add(x.get(story_key), x.get(drift_key), x.get(case_key, case) if case_key else case)
                return
            for key, val in x.items():
                next_case = case
                if isinstance(key, str) and not key.strip().replace(".", "", 1).isdigit():
                    next_case = key
                if isinstance(val, (dict, list, tuple, pd.DataFrame)):
                    walk(val, next_case)
                else:
                    add(key, val, next_case)
            return
        if isinstance(x, (list, tuple, set)):
            for item in x:
                walk(item, case)
            return

    walk(obj, default_case)
    out = pd.DataFrame(rows)
    if out.empty:
        return pd.DataFrame(columns=["Story", "Case", "Drift_Ratio"])
    out["Story"] = pd.to_numeric(out["Story"], errors="coerce")
    out["Drift_Ratio"] = pd.to_numeric(out["Drift_Ratio"], errors="coerce").abs()
    out = out.dropna(subset=["Story", "Drift_Ratio"])
    out = out[out["Story"].astype(str).str.lower() != "envelope"]
    out["Story"] = out["Story"].astype(int)
    return out[["Story", "Case", "Drift_Ratio"]]


def _envelope_by_story(rows: pd.DataFrame, prefix: str) -> pd.DataFrame:
    if rows is None or rows.empty:
        return pd.DataFrame(columns=["Story", f"Governing_Case_{prefix}", f"Drift_{prefix}_Ratio"])
    d = rows.copy()
    d["Story"] = pd.to_numeric(d["Story"], errors="coerce")
    d["Drift_Ratio"] = pd.to_numeric(d["Drift_Ratio"], errors="coerce").abs()
    d = d.dropna(subset=["Story", "Drift_Ratio"])
    if d.empty:
        return pd.DataFrame(columns=["Story", f"Governing_Case_{prefix}", f"Drift_{prefix}_Ratio"])
    idx = d.groupby("Story")["Drift_Ratio"].idxmax()
    out = d.loc[idx, ["Story", "Case", "Drift_Ratio"]].copy()
    out["Story"] = out["Story"].astype(int)
    out = out.sort_values("Story")
    return out.rename(columns={"Case": f"Governing_Case_{prefix}", "Drift_Ratio": f"Drift_{prefix}_Ratio"})


def _selected_wall_drift_rows(sw_opt_df: Optional[pd.DataFrame]) -> pd.DataFrame:
    if sw_opt_df is None or getattr(sw_opt_df, "empty", True):
        return pd.DataFrame(columns=["Story", "Case", "Drift_Ratio"])
    r = sw_opt_df.iloc[0]
    for key in ["_drifts", "Selected_Drifts", "drifts", "Drifts"]:
        if key in sw_opt_df.columns:
            rows = drift_rows_from_any(r.get(key), default_case="After selected shear walls")
            if not rows.empty:
                return rows
    return pd.DataFrame(columns=["Story", "Case", "Drift_Ratio"])


def build_story_drift_comparison(
    df_static_drifts_before: Optional[pd.DataFrame],
    sw_opt_df: Optional[pd.DataFrame] = None,
) -> pd.DataFrame:
    before = _envelope_by_story(drift_rows_from_any(df_static_drifts_before, "Before retrofit"), "Before")
    after = _envelope_by_story(_selected_wall_drift_rows(sw_opt_df), "After")

    if before.empty and after.empty:
        return pd.DataFrame()
    out = before.merge(after, on="Story", how="outer").sort_values("Story").reset_index(drop=True)
    if "Drift_Before_Ratio" in out.columns and "Drift_After_Ratio" in out.columns:
        out["Drift_Reduction_pct"] = 100.0 * (out["Drift_Before_Ratio"] - out["Drift_After_Ratio"]) / out["Drift_Before_Ratio"].replace(0, np.nan)
    return out


# ============================================================
# Member comparison and summaries
# ============================================================


def _extract_member_eta(df: Optional[pd.DataFrame], label: str) -> pd.DataFrame:
    """Extract total EC2 utilization and, when available, shear/flexure components."""
    if df is None or getattr(df, "empty", True):
        return pd.DataFrame(columns=["Element", label])

    ele_col = _first_col(df, ["Element", "ElementID", "Member", "Member ID"])
    eta_col = _first_col(df, ["eta", "Static_EC2_Utilization", "Utilization", "Worst eta", "Max_Eta"])
    type_col = _first_col(df, ["Type", "etype", "Member_Type"])
    eta_v_col = _first_col(df, ["etaV", "EtaV", "eta_shear", "Shear_Utilization", "V_Utilization"])
    eta_nm_col = _first_col(df, ["eta_NM", "etaNM", "Eta_NM", "eta_flexure", "Flexure_Axial_Utilization", "NM_Utilization"])

    if ele_col is None or eta_col is None:
        return pd.DataFrame(columns=["Element", label])

    out = pd.DataFrame({
        "Element": pd.to_numeric(df[ele_col], errors="coerce"),
        label: pd.to_numeric(df[eta_col], errors="coerce"),
    })
    if type_col is not None:
        out["Type"] = df[type_col].astype(str).values
    if eta_v_col is not None:
        out[f"{label}_Shear"] = pd.to_numeric(df[eta_v_col], errors="coerce")
    if eta_nm_col is not None:
        out[f"{label}_Flexure_Axial"] = pd.to_numeric(df[eta_nm_col], errors="coerce")

    out = out.dropna(subset=["Element"]).copy()
    out["Element"] = out["Element"].astype(int)
    return out.drop_duplicates("Element")


def _diagnose_problem_from_components(row: pd.Series, shear_col: str, flex_col: str, total_col: str) -> str:
    shear = _safe_float(row.get(shear_col, np.nan))
    flex = _safe_float(row.get(flex_col, np.nan))
    total = _safe_float(row.get(total_col, np.nan))
    if np.isfinite(shear) and np.isfinite(flex):
        if shear >= flex and shear > 1.0:
            return "SHEAR"
        if flex > shear and flex > 1.0:
            return "FLEXURE/AXIAL"
        return "OK" if np.isfinite(total) and total <= 1.0 else ("SHEAR" if shear >= flex else "FLEXURE/AXIAL")
    if np.isfinite(total):
        return "OK" if total <= 1.0 else "UNKNOWN"
    return "UNKNOWN"


def _normalize_jacketing_member_report(jacketing_member_report: Optional[pd.DataFrame]) -> pd.DataFrame:
    """Normalize the optimizer's per-member decision report for Excel/Word export."""
    if jacketing_member_report is None or getattr(jacketing_member_report, "empty", True):
        return pd.DataFrame()

    rpt = jacketing_member_report.copy()
    ele_col = _first_col(rpt, ["Element", "ElementID", "Member", "Member ID"])
    if ele_col is None:
        return pd.DataFrame()
    rpt["Element"] = pd.to_numeric(rpt[ele_col], errors="coerce")
    rpt = rpt.dropna(subset=["Element"]).copy()
    rpt["Element"] = rpt["Element"].astype(int)

    rename = {
        "Original_Eta": "Eta_Before_Optimizer",
        "Final_Eta": "Eta_Final_Optimizer",
        "Original_Eta_Shear": "Eta_Before_Shear_Optimizer",
        "Original_Eta_Flexure_Axial": "Eta_Before_Flexure_Axial_Optimizer",
        "Final_Eta_Shear": "Eta_Final_Shear_Optimizer",
        "Final_Eta_Flexure_Axial": "Eta_Final_Flexure_Axial_Optimizer",
    }
    rpt = rpt.rename(columns={k: v for k, v in rename.items() if k in rpt.columns})

    wanted = [
        "Element", "Problem_Before", "Eta_Before_Optimizer", "Eta_Before_Shear_Optimizer",
        "Eta_Before_Flexure_Axial_Optimizer", "Problem_After", "Eta_Final_Optimizer",
        "Eta_Final_Shear_Optimizer", "Eta_Final_Flexure_Axial_Optimizer",
        "Chosen_Jacket_Thickness_mm", "Chosen_Stirrup_Spacing_mm",
        "Chosen_Longitudinal_Rebar_Increase_pct",
        "Final_Rebar_Min_Clear_Spacing_mm", "Final_Rebar_Required_Clear_Spacing_mm",
        "Final_Rebar_Spacing_Status", "Final_Rebar_Spacing_Critical_Face",
        "Final_Rebar_Spacing_Note", "Final_Status",
    ]
    return rpt[[c for c in wanted if c in rpt.columns]].drop_duplicates("Element")


def build_member_utilization_comparison(
    df_worst_before: Optional[pd.DataFrame],
    df_worst_after_walls: Optional[pd.DataFrame],
    df_worst_final: Optional[pd.DataFrame],
    jacketing_member_report: Optional[pd.DataFrame] = None,
    global_opt_df: Optional[pd.DataFrame] = None,
) -> pd.DataFrame:
    """Build final member table, including shear/flexure diagnosis and jacketing decisions."""
    base = _extract_member_eta(df_worst_before, "Eta_Before")
    walls = _extract_member_eta(df_worst_after_walls, "Eta_After_Walls")
    final = _extract_member_eta(df_worst_final if df_worst_final is not None else df_worst_after_walls, "Eta_Final")
    decision = _normalize_jacketing_member_report(jacketing_member_report)

    ids = sorted(
        set(base.get("Element", []))
        .union(walls.get("Element", []))
        .union(final.get("Element", []))
        .union(decision.get("Element", []))
    )
    if not ids:
        return pd.DataFrame()

    out = pd.DataFrame({"Element": ids})
    for frame in [base, walls, final, decision]:
        if frame is not None and not frame.empty:
            out = out.merge(frame, on="Element", how="left")

    type_cols = [c for c in out.columns if c.startswith("Type")]
    if type_cols:
        out["Type"] = ""
        for c in type_cols:
            out["Type"] = np.where(out["Type"].astype(str).str.len() > 0, out["Type"], out[c].fillna(""))
        out = out.drop(columns=type_cols)

    def _col_series(name: str) -> pd.Series:
        if name in out.columns:
            return pd.to_numeric(out[name], errors="coerce")
        return pd.Series(np.nan, index=out.index)

    # Prefer the optimizer report values for jacketed members because they are
    # taken directly from the selected automatic-jacketing iteration.
    if "Eta_Before_Optimizer" in out.columns:
        out["Eta_Before"] = _col_series("Eta_Before").where(_col_series("Eta_Before").notna(), _col_series("Eta_Before_Optimizer"))
    if "Eta_Final_Optimizer" in out.columns:
        out["Eta_Final"] = _col_series("Eta_Final").where(_col_series("Eta_Final").notna(), _col_series("Eta_Final_Optimizer"))
    if "Eta_Before_Shear_Optimizer" in out.columns:
        out["Eta_Before_Shear"] = _col_series("Eta_Before_Shear").where(_col_series("Eta_Before_Shear").notna(), _col_series("Eta_Before_Shear_Optimizer"))
    if "Eta_Before_Flexure_Axial_Optimizer" in out.columns:
        out["Eta_Before_Flexure_Axial"] = _col_series("Eta_Before_Flexure_Axial").where(_col_series("Eta_Before_Flexure_Axial").notna(), _col_series("Eta_Before_Flexure_Axial_Optimizer"))
    if "Eta_Final_Shear_Optimizer" in out.columns:
        out["Eta_Final_Shear"] = _col_series("Eta_Final_Shear").where(_col_series("Eta_Final_Shear").notna(), _col_series("Eta_Final_Shear_Optimizer"))
    if "Eta_Final_Flexure_Axial_Optimizer" in out.columns:
        out["Eta_Final_Flexure_Axial"] = _col_series("Eta_Final_Flexure_Axial").where(_col_series("Eta_Final_Flexure_Axial").notna(), _col_series("Eta_Final_Flexure_Axial_Optimizer"))

    if "Problem_Before" not in out.columns:
        out["Problem_Before"] = out.apply(lambda r: _diagnose_problem_from_components(r, "Eta_Before_Shear", "Eta_Before_Flexure_Axial", "Eta_Before"), axis=1)
    else:
        out["Problem_Before"] = out["Problem_Before"].fillna(out.apply(lambda r: _diagnose_problem_from_components(r, "Eta_Before_Shear", "Eta_Before_Flexure_Axial", "Eta_Before"), axis=1))

    if "Problem_After" not in out.columns:
        out["Problem_After"] = out.apply(lambda r: _diagnose_problem_from_components(r, "Eta_Final_Shear", "Eta_Final_Flexure_Axial", "Eta_Final"), axis=1)
    else:
        out["Problem_After"] = out["Problem_After"].fillna(out.apply(lambda r: _diagnose_problem_from_components(r, "Eta_Final_Shear", "Eta_Final_Flexure_Axial", "Eta_Final"), axis=1))

    if "Eta_Before" in out.columns and "Eta_Final" in out.columns:
        out["Eta_Reduction_pct"] = 100.0 * (out["Eta_Before"] - out["Eta_Final"]) / out["Eta_Before"].replace(0, np.nan)
    if "Final_Status" not in out.columns and "Eta_Final" in out.columns:
        out["Final_Status"] = out["Eta_Final"].apply(lambda v: "OK" if pd.notna(v) and float(v) <= 1.0 else "FAIL")
    elif "Final_Status" in out.columns and "Eta_Final" in out.columns:
        inferred = out["Eta_Final"].apply(lambda v: "OK" if pd.notna(v) and float(v) <= 1.0 else "FAIL")
        out["Final_Status"] = out["Final_Status"].fillna(inferred)

    drop_cols = [c for c in out.columns if c.endswith("_Optimizer")]
    out = out.drop(columns=drop_cols, errors="ignore")

    cols = ["Element"] + (["Type"] if "Type" in out.columns else [])
    cols += [c for c in [
        "Eta_Before", "Eta_Before_Shear", "Eta_Before_Flexure_Axial", "Problem_Before",
        "Eta_After_Walls", "Eta_After_Walls_Shear", "Eta_After_Walls_Flexure_Axial",
        "Eta_Final", "Eta_Final_Shear", "Eta_Final_Flexure_Axial", "Problem_After",
        "Eta_Reduction_pct", "Chosen_Jacket_Thickness_mm", "Chosen_Stirrup_Spacing_mm",
        "Chosen_Longitudinal_Rebar_Increase_pct",
        "Final_Rebar_Min_Clear_Spacing_mm", "Final_Rebar_Required_Clear_Spacing_mm",
        "Final_Rebar_Spacing_Status", "Final_Rebar_Spacing_Critical_Face",
        "Final_Rebar_Spacing_Note", "Final_Status",
    ] if c in out.columns]
    return out[cols]

def _count_failing_members(df: Optional[pd.DataFrame]) -> int:
    if df is None or getattr(df, "empty", True):
        return 0
    for col in ["Final_Status", "Status", "Pass", "StaticPass", "Overall_Static_Status"]:
        real = _first_col(df, [col])
        if real is not None:
            txt = df[real].astype(str).str.upper()
            return int(txt.isin(["FAIL", "NO", "FALSE"]).sum())
    eta = _max_from_cols(df, ["eta", "Static_EC2_Utilization", "Utilization", "Worst eta", "Max_Eta"])
    if np.isfinite(eta):
        real = _first_col(df, ["eta", "Static_EC2_Utilization", "Utilization", "Worst eta", "Max_Eta"])
        return int((pd.to_numeric(df[real], errors="coerce") > 1.0).sum())
    return 0


def build_executive_summary(
    *,
    drift_compare_df: Optional[pd.DataFrame],
    drift_limit_ratio: Optional[float],
    sw_opt_df: Optional[pd.DataFrame],
    df_worst_before: Optional[pd.DataFrame],
    df_worst_after_walls: Optional[pd.DataFrame],
    df_worst_final: Optional[pd.DataFrame],
    jacketing_trials_df: Optional[pd.DataFrame],
    jacket_cost: Any,
    jacketing_member_report: Optional[pd.DataFrame] = None,
    global_opt_df: Optional[pd.DataFrame] = None,
) -> pd.DataFrame:
    before_drift = _max_from_cols(drift_compare_df, ["Drift_Before_Ratio"])
    after_walls_drift = _max_from_cols(drift_compare_df, ["Drift_After_Ratio"])
    if not np.isfinite(after_walls_drift) and sw_opt_df is not None and not sw_opt_df.empty:
        after_walls_drift = _safe_float(sw_opt_df.iloc[0].get("Max_Drift_Ratio", np.nan))
    reduction = np.nan
    if np.isfinite(before_drift) and abs(before_drift) > 1e-12 and np.isfinite(after_walls_drift):
        reduction = 100.0 * (before_drift - after_walls_drift) / before_drift

    wall_cost = np.nan
    wall_desc = ""
    if sw_opt_df is not None and not sw_opt_df.empty:
        wall_cost = _safe_float(sw_opt_df.iloc[0].get("Total_Cost", np.nan))
        wall_desc = str(sw_opt_df.iloc[0].get("Description", ""))
    jacket_total = _safe_float(_cost_to_dict(jacket_cost).get("total_cost", np.nan))
    if not np.isfinite(jacket_total) and jacketing_trials_df is not None and not jacketing_trials_df.empty:
        jacket_total = _max_from_cols(jacketing_trials_df, ["Total_Cost"])

    member_diag_df = _normalize_jacketing_member_report(jacketing_member_report)
    n_shear_before = 0
    n_flex_before = 0
    n_shear_after = 0
    n_flex_after = 0
    if not member_diag_df.empty:
        if "Problem_Before" in member_diag_df.columns:
            before_txt = member_diag_df["Problem_Before"].astype(str).str.upper()
            n_shear_before = int(before_txt.str.contains("SHEAR", na=False).sum())
            n_flex_before = int((before_txt.str.contains("FLEX", na=False) | before_txt.str.contains("MOMENT", na=False) | before_txt.str.contains("AXIAL", na=False)).sum())
        if "Problem_After" in member_diag_df.columns:
            still_fail = member_diag_df["Final_Status"].astype(str).str.upper().isin(["FAIL"]) if "Final_Status" in member_diag_df.columns else True
            after_txt = member_diag_df["Problem_After"].astype(str).str.upper().where(still_fail, "")
            n_shear_after = int(after_txt.str.contains("SHEAR", na=False).sum())
            n_flex_after = int((after_txt.str.contains("FLEX", na=False) | after_txt.str.contains("MOMENT", na=False) | after_txt.str.contains("AXIAL", na=False)).sum())

    rows = [
        ["Drift limit ratio", drift_limit_ratio],
        ["Maximum drift before retrofit", before_drift],
        ["Maximum drift after selected shear walls", after_walls_drift],
        ["Drift reduction after selected shear walls [%]", reduction],
        ["Drift status after selected shear walls", "OK" if np.isfinite(after_walls_drift) and drift_limit_ratio is not None and after_walls_drift <= drift_limit_ratio else "FAIL"],
        ["Selected wall layout", wall_desc],
        ["Estimated shear-wall cost", wall_cost],
        ["Selected wall objective score", _safe_float(sw_opt_df.iloc[0].get("Objective_Total_Score", np.nan)) if sw_opt_df is not None and not sw_opt_df.empty else np.nan],
        ["Selected wall objective rank", _safe_float(sw_opt_df.iloc[0].get("Objective_Rank", np.nan)) if sw_opt_df is not None and not sw_opt_df.empty else np.nan],
        ["Selected wall constructability class", str(sw_opt_df.iloc[0].get("Constructability_Class", "")) if sw_opt_df is not None and not sw_opt_df.empty else ""],
        ["Members failing before retrofit", _count_failing_members(df_worst_before)],
        ["Members failing after selected shear walls", _count_failing_members(df_worst_after_walls)],
        ["Members failing after final jacketing", _count_failing_members(df_worst_final if df_worst_final is not None else df_worst_after_walls)],
        ["Final member status", "OK" if _count_failing_members(df_worst_final if df_worst_final is not None else df_worst_after_walls) == 0 else "FAIL"],
        ["Shear-governed jacketed members before jacketing", n_shear_before],
        ["Flexure/axial-governed jacketed members before jacketing", n_flex_before],
        ["Shear-governed jacketed members still critical after jacketing", n_shear_after],
        ["Flexure/axial-governed jacketed members still critical after jacketing", n_flex_after],
        ["Estimated concrete jacketing cost", jacket_total],
        ["Total estimated retrofit cost", _safe_float(wall_cost, 0.0) + _safe_float(jacket_total, 0.0)],
    ]

    if global_opt_df is not None and not getattr(global_opt_df, "empty", True):
        try:
            g = global_opt_df.iloc[0]
            rows += [
                ["Global optimizer selected candidate", str(g.get("Candidate_ID", ""))],
                ["Global optimizer feasibility", str(g.get("Complete_Feasible", ""))],
                ["Global search candidate limit", str(g.get("Search_Max_Global_Candidates", ""))],
                ["Global candidates sent to expensive stage", _safe_float(g.get("Search_Candidates_Sent_To_Global_Stage", np.nan))],
                ["Global unique complete strategies", _safe_float(g.get("Search_Unique_Complete_Strategies", np.nan))],
                ["Global total retrofit cost", _safe_float(g.get("Total_Retrofit_Cost", np.nan))],
                ["Global wall cost", _safe_float(g.get("Wall_Cost", np.nan))],
                ["Global jacketing cost", _safe_float(g.get("Jacketing_Cost", np.nan))],
                ["Global final max EC2 utilization", _safe_float(g.get("Final_Max_EC2_Utilization", np.nan))],
                ["Global final failed members", _safe_float(g.get("Final_Failed_Members", np.nan))],
            ]
        except Exception:
            pass
    return pd.DataFrame(rows, columns=["Metric", "Value"])


# ============================================================
# Failed-member diagnosis and N/M/V demand diagrams
# ============================================================


def _extract_member_demand_snapshot(df: Optional[pd.DataFrame], prefix: str) -> pd.DataFrame:
    """Extract EC2 utilization components and N/V/M actions per element.

    The EC2 worst table normally stores one critical row per member. These
    values are used to create a professional demand snapshot in the report. If
    full station-by-station diagram data are added later, this function can be
    extended without changing the public report API.
    """
    if df is None or getattr(df, "empty", True):
        return pd.DataFrame(columns=["Element"])

    ele_col = _first_col(df, ["Element", "ElementID", "Member", "Member ID"])
    if ele_col is None:
        return pd.DataFrame(columns=["Element"])

    type_col = _first_col(df, ["Type", "etype", "Member_Type"])
    combo_col = _first_col(df, ["Combination", "Combo", "Case", "Load_Combination"])
    eta_col = _first_col(df, ["eta", "Static_EC2_Utilization", "Utilization", "Worst eta", "Max_Eta"])
    eta_v_col = _first_col(df, ["etaV", "EtaV", "eta_shear", "Shear_Utilization", "V_Utilization"])
    eta_nm_col = _first_col(df, ["eta_NM", "etaNM", "Eta_NM", "eta_flexure", "Flexure_Axial_Utilization", "NM_Utilization"])
    n_col = _first_col(df, ["NEd_kN", "N_kN", "N", "Axial_kN", "AxialForce_kN"])
    v_col = _first_col(df, ["VEd_kN", "V_kN", "V", "Shear_kN", "ShearForce_kN"])
    m_col = _first_col(df, ["MEd_kNm", "M_kNm", "M", "Moment_kNm", "BendingMoment_kNm"])

    rows = []
    for _, rr in df.iterrows():
        try:
            eid = int(pd.to_numeric(rr[ele_col], errors="coerce"))
        except Exception:
            continue
        row = {"Element": eid}
        if type_col is not None:
            row[f"Type_{prefix}"] = str(rr.get(type_col, ""))
        if combo_col is not None:
            row[f"Combo_{prefix}"] = str(rr.get(combo_col, ""))
        if eta_col is not None:
            row[f"Eta_{prefix}"] = _safe_float(rr.get(eta_col, np.nan))
        if eta_v_col is not None:
            row[f"Eta_Shear_{prefix}"] = _safe_float(rr.get(eta_v_col, np.nan))
        if eta_nm_col is not None:
            row[f"Eta_Flexure_Axial_{prefix}"] = _safe_float(rr.get(eta_nm_col, np.nan))
        if n_col is not None:
            row[f"NEd_kN_{prefix}"] = _safe_float(rr.get(n_col, np.nan))
        if v_col is not None:
            row[f"VEd_kN_{prefix}"] = _safe_float(rr.get(v_col, np.nan))
        if m_col is not None:
            row[f"MEd_kNm_{prefix}"] = _safe_float(rr.get(m_col, np.nan))
        rows.append(row)

    out = pd.DataFrame(rows)
    if out.empty:
        return pd.DataFrame(columns=["Element"])
    # Keep the largest utilization row if duplicate elements are present.
    eta_name = f"Eta_{prefix}"
    if eta_name in out.columns:
        out[eta_name] = pd.to_numeric(out[eta_name], errors="coerce")
        out = out.sort_values(eta_name, ascending=False)
    return out.drop_duplicates("Element")


def _status_from_eta(value: Any) -> str:
    eta = _safe_float(value, np.nan)
    if not np.isfinite(eta):
        return "UNKNOWN"
    return "OK" if eta <= 1.0 else "FAIL"


def _recommendation_from_problem(problem: str, eta_shear: float, eta_flex: float) -> str:
    p = str(problem or "").upper()
    if "SHEAR" in p:
        return "Reduce stirrup spacing / increase transverse reinforcement; verify shear capacity and detailing limits."
    if "FLEX" in p or "MOMENT" in p or "AXIAL" in p:
        return "Increase longitudinal reinforcement first; use jacket thickness increase only if reinforcement/detailing limits are reached."
    if np.isfinite(eta_shear) and np.isfinite(eta_flex):
        if eta_shear >= eta_flex and eta_shear > 1.0:
            return "Shear demand controls; prioritize transverse reinforcement."
        if eta_flex > eta_shear and eta_flex > 1.0:
            return "Flexure/axial demand controls; prioritize longitudinal reinforcement."
    return "Review detailed EC2 checks; governing failure mode could not be isolated from available columns."


def build_failed_member_diagnosis_table(
    *,
    member_compare_df: Optional[pd.DataFrame],
    df_worst_after_walls: Optional[pd.DataFrame],
    df_worst_final: Optional[pd.DataFrame],
    member_fail_after_walls: Optional[pd.DataFrame],
    member_fail_final: Optional[pd.DataFrame],
) -> pd.DataFrame:
    """Create a focused diagnosis table for members that failed at any retrofit stage."""
    frames = []
    for df in [member_compare_df, member_fail_after_walls, member_fail_final, df_worst_after_walls, df_worst_final]:
        if df is None or getattr(df, "empty", True):
            continue
        ele_col = _first_col(df, ["Element", "ElementID", "Member", "Member ID"])
        if ele_col is None:
            continue
        tmp = pd.DataFrame({"Element": pd.to_numeric(df[ele_col], errors="coerce")}).dropna()
        if not tmp.empty:
            frames.append(tmp)
    if not frames:
        return pd.DataFrame()

    ids_df = pd.concat(frames, ignore_index=True)
    ids_df["Element"] = ids_df["Element"].astype(int)
    out = pd.DataFrame({"Element": sorted(ids_df["Element"].unique())})

    after_walls = _extract_member_demand_snapshot(df_worst_after_walls, "After_Walls")
    final = _extract_member_demand_snapshot(df_worst_final, "Final")
    for frame in [after_walls, final]:
        if frame is not None and not frame.empty:
            out = out.merge(frame, on="Element", how="left")

    if member_compare_df is not None and not getattr(member_compare_df, "empty", True):
        comp = member_compare_df.copy()
        ele_col = _first_col(comp, ["Element", "ElementID", "Member", "Member ID"])
        if ele_col is not None:
            comp["Element"] = pd.to_numeric(comp[ele_col], errors="coerce")
            comp = comp.dropna(subset=["Element"]).copy()
            comp["Element"] = comp["Element"].astype(int)
            keep = [
                "Element", "Type", "Eta_After_Walls", "Eta_After_Walls_Shear", "Eta_After_Walls_Flexure_Axial",
                "Eta_Final", "Eta_Final_Shear", "Eta_Final_Flexure_Axial", "Problem_After", "Problem_Before",
                "Chosen_Jacket_Thickness_mm", "Chosen_Stirrup_Spacing_mm", "Chosen_Longitudinal_Rebar_Increase_pct",
                "Final_Rebar_Min_Clear_Spacing_mm", "Final_Rebar_Required_Clear_Spacing_mm",
                "Final_Rebar_Spacing_Status", "Final_Rebar_Spacing_Critical_Face",
                "Final_Rebar_Spacing_Note", "Final_Status",
            ]
            comp = comp[[c for c in keep if c in comp.columns]].drop_duplicates("Element")
            out = out.merge(comp, on="Element", how="left", suffixes=("", "_Report"))

    # Consolidate type and eta/problem columns.
    type_cols = [c for c in ["Type", "Type_Final", "Type_After_Walls"] if c in out.columns]
    if type_cols:
        out["Member_Type"] = ""
        for c in type_cols:
            out["Member_Type"] = np.where(out["Member_Type"].astype(str).str.len() > 0, out["Member_Type"], out[c].fillna(""))

    for target, candidates in {
        "Eta_After_Walls": ["Eta_After_Walls", "Eta_After_Walls_Report"],
        "Eta_Shear_After_Walls": ["Eta_Shear_After_Walls", "Eta_After_Walls_Shear"],
        "Eta_Flexure_Axial_After_Walls": ["Eta_Flexure_Axial_After_Walls", "Eta_After_Walls_Flexure_Axial"],
        "Eta_Final": ["Eta_Final", "Eta_Final_Report"],
        "Eta_Shear_Final": ["Eta_Shear_Final", "Eta_Final_Shear"],
        "Eta_Flexure_Axial_Final": ["Eta_Flexure_Axial_Final", "Eta_Final_Flexure_Axial"],
    }.items():
        vals = pd.Series(np.nan, index=out.index)
        for c in candidates:
            if c in out.columns:
                vals = vals.where(vals.notna(), pd.to_numeric(out[c], errors="coerce"))
        out[target] = vals

    def diagnose_row(rr: pd.Series, stage: str) -> str:
        problem_col = "Problem_After" if stage == "Final" else "Problem_Before"
        if stage == "Final" and problem_col in rr and str(rr.get(problem_col, "")).strip() and str(rr.get(problem_col, "")).upper() != "NAN":
            return str(rr.get(problem_col))
        return _diagnose_problem_from_components(
            rr,
            f"Eta_Shear_{stage}",
            f"Eta_Flexure_Axial_{stage}",
            f"Eta_{stage}",
        )

    out["Problem_After_Walls"] = out.apply(lambda r: _diagnose_problem_from_components(r, "Eta_Shear_After_Walls", "Eta_Flexure_Axial_After_Walls", "Eta_After_Walls"), axis=1)
    out["Problem_Final"] = out.apply(lambda r: diagnose_row(r, "Final"), axis=1)
    out["Status_After_Walls"] = out["Eta_After_Walls"].apply(_status_from_eta)
    out["Status_Final"] = out.get("Final_Status", pd.Series(index=out.index, dtype=object)).fillna(out["Eta_Final"].apply(_status_from_eta))

    out["Engineering_Recommendation"] = out.apply(
        lambda r: _recommendation_from_problem(
            r.get("Problem_Final") if str(r.get("Status_Final", "")).upper() == "FAIL" else r.get("Problem_After_Walls"),
            _safe_float(r.get("Eta_Shear_Final", r.get("Eta_Shear_After_Walls", np.nan))),
            _safe_float(r.get("Eta_Flexure_Axial_Final", r.get("Eta_Flexure_Axial_After_Walls", np.nan))),
        ),
        axis=1,
    )

    # Focus the table on members that are or were critical; keep jacketed members even if final OK.
    critical_mask = (
        out["Status_After_Walls"].astype(str).str.upper().eq("FAIL")
        | out["Status_Final"].astype(str).str.upper().eq("FAIL")
        | out[[c for c in ["Chosen_Jacket_Thickness_mm", "Chosen_Stirrup_Spacing_mm", "Chosen_Longitudinal_Rebar_Increase_pct"] if c in out.columns]].notna().any(axis=1)
        | (out["Final_Rebar_Spacing_Status"].astype(str).str.upper().eq("FAIL") if "Final_Rebar_Spacing_Status" in out.columns else False)
    )
    out = out.loc[critical_mask].copy()

    wanted = [
        "Element", "Member_Type", "Combo_After_Walls", "Combo_Final",
        "NEd_kN_After_Walls", "VEd_kN_After_Walls", "MEd_kNm_After_Walls",
        "Eta_After_Walls", "Eta_Shear_After_Walls", "Eta_Flexure_Axial_After_Walls", "Problem_After_Walls", "Status_After_Walls",
        "NEd_kN_Final", "VEd_kN_Final", "MEd_kNm_Final",
        "Eta_Final", "Eta_Shear_Final", "Eta_Flexure_Axial_Final", "Problem_Final", "Status_Final",
        "Chosen_Jacket_Thickness_mm", "Chosen_Stirrup_Spacing_mm", "Chosen_Longitudinal_Rebar_Increase_pct",
        "Final_Rebar_Min_Clear_Spacing_mm", "Final_Rebar_Required_Clear_Spacing_mm",
        "Final_Rebar_Spacing_Status", "Final_Rebar_Spacing_Critical_Face",
        "Final_Rebar_Spacing_Note",
        "Engineering_Recommendation",
    ]
    out = out[[c for c in wanted if c in out.columns]]
    sort_col = "Eta_Final" if "Eta_Final" in out.columns else "Eta_After_Walls"
    if sort_col in out.columns:
        out[sort_col] = pd.to_numeric(out[sort_col], errors="coerce")
        out = out.sort_values(["Status_Final", sort_col], ascending=[True, False], na_position="last")
    return out.reset_index(drop=True)


def _diagram_values_from_row(row: pd.Series, stage: str) -> tuple[float, float, float]:
    n = _safe_float(row.get(f"NEd_kN_{stage}", np.nan), 0.0)
    v = _safe_float(row.get(f"VEd_kN_{stage}", np.nan), 0.0)
    m = _safe_float(row.get(f"MEd_kNm_{stage}", np.nan), 0.0)
    return n, v, m


def _member_length_m_for_report(row: pd.Series, nodes: Optional[dict], elements: Optional[dict]) -> float:
    """Return the real member length in metres when nodes/elements are available."""
    try:
        eid = int(row.get("Element"))
        ni, nj, _etype = elements[eid]
        x1, y1 = nodes[ni]
        x2, y2 = nodes[nj]
        L = float(np.hypot(float(x2) - float(x1), float(y2) - float(y1)))
        if np.isfinite(L) and L > 0.0:
            return L
    except Exception:
        pass
    return 1.0


def _normalize_force_diagram_data(data: Any) -> Optional[pd.DataFrame]:
    """Normalize optional station-by-station N/V/M data to x_m, N_kN, V_kN, M_kNm.

    Accepted forms are dataframes or dictionaries with columns/keys similar to
    x, station, N, V and M. This keeps the report ready for exact OpenSees
    envelopes when the analysis pipeline starts passing them in.
    """
    if data is None:
        return None
    try:
        df = data.copy() if isinstance(data, pd.DataFrame) else pd.DataFrame(data)
    except Exception:
        return None
    if df.empty:
        return None

    x_col = _first_col(df, ["x_m", "x", "station_m", "station", "Distance_m", "Distance", "s_m"])
    n_col = _first_col(df, ["N_kN", "N", "NEd_kN", "axial_kN", "Axial_kN"])
    v_col = _first_col(df, ["V_kN", "V", "VEd_kN", "shear_kN", "Shear_kN"])
    m_col = _first_col(df, ["M_kNm", "M", "MEd_kNm", "moment_kNm", "Moment_kNm"])
    if x_col is None or (n_col is None and v_col is None and m_col is None):
        return None

    out = pd.DataFrame({"x_m": pd.to_numeric(df[x_col], errors="coerce")})
    out["N_kN"] = pd.to_numeric(df[n_col], errors="coerce") if n_col is not None else np.nan
    out["V_kN"] = pd.to_numeric(df[v_col], errors="coerce") if v_col is not None else np.nan
    out["M_kNm"] = pd.to_numeric(df[m_col], errors="coerce") if m_col is not None else np.nan
    out = out.dropna(subset=["x_m"]).sort_values("x_m")
    if out.empty:
        return None
    return out


def _force_diagram_for_member(member_force_diagrams: Any, element: int, stage: str) -> Optional[pd.DataFrame]:
    """Extract optional exact N/V/M station data for a member and stage."""
    if member_force_diagrams is None:
        return None
    candidates = []
    for key in (element, str(element)):
        try:
            candidates.append(member_force_diagrams.get(key))
        except Exception:
            pass
    for obj in candidates:
        if obj is None:
            continue
        if isinstance(obj, dict):
            for stage_key in (stage, stage.lower(), stage.replace("_", " "), stage.replace("_", "").lower()):
                if stage_key in obj:
                    norm = _normalize_force_diagram_data(obj[stage_key])
                    if norm is not None:
                        return norm
        norm = _normalize_force_diagram_data(obj)
        if norm is not None:
            return norm
    return None


def _plot_single_member_nmv(
    row: pd.Series,
    filename: str,
    *,
    nodes: Optional[dict] = None,
    elements: Optional[dict] = None,
    member_force_diagrams: Any = None,
) -> Optional[str]:
    """Save N/V/M plots for one critical member.

    Exact station-by-station N/V/M diagrams are plotted when they are supplied
    through member_force_diagrams. If only EC2 worst-result ordinates are
    available, the report now plots them as constant critical-demand ordinates
    along the real member length. It no longer invents a parabolic moment shape
    or zero end moments, because that is not valid for general distributed-load
    frame members.
    """
    element = row.get("Element", "?")
    try:
        element_int = int(element)
    except Exception:
        element_int = -1
    length_m = _member_length_m_for_report(row, nodes, elements)

    stages = []
    used_exact_data = False
    for stage, label in [("After_Walls", "After walls"), ("Final", "Final")]:
        exact = _force_diagram_for_member(member_force_diagrams, element_int, stage)
        if exact is not None:
            stages.append((label, exact))
            used_exact_data = True
            continue
        n, v, m = _diagram_values_from_row(row, stage)
        if any(np.isfinite(x) and abs(x) > 1e-12 for x in [n, v, m]):
            x = np.linspace(0.0, length_m, 25)
            stages.append((label, pd.DataFrame({
                "x_m": x,
                "N_kN": np.full_like(x, n),
                "V_kN": np.full_like(x, v),
                "M_kNm": np.full_like(x, m),
            })))
    if not stages:
        return None

    fig, axes = plt.subplots(3, 1, figsize=(7.6, 6.5), sharex=True)
    for label, data in stages:
        x = pd.to_numeric(data["x_m"], errors="coerce")
        axes[0].plot(x, pd.to_numeric(data["N_kN"], errors="coerce"), label=label)
        axes[1].plot(x, pd.to_numeric(data["V_kN"], errors="coerce"), label=label)
        axes[2].plot(x, pd.to_numeric(data["M_kNm"], errors="coerce"), label=label)

    axes[0].set_ylabel("N [kN]")
    axes[1].set_ylabel("V [kN]")
    axes[2].set_ylabel("M [kNm]  (sagging + downward)")
    axes[2].invert_yaxis()      # moments on the tension side
    axes[2].set_xlabel("Member coordinate x [m]")
    for ax in axes:
        ax.axhline(0.0, linewidth=0.8)
        ax.grid(True, alpha=0.3)
        ax.legend(fontsize=8)
    axes[2].set_xlim(0.0, max(length_m, 1e-6))

    problem = str(row.get("Problem_Final", row.get("Problem_After_Walls", "")))
    status = str(row.get("Status_Final", ""))
    diagram_kind = "station-by-station N/V/M diagram" if used_exact_data else "critical EC2 demand ordinates"
    fig.suptitle(
        f"Critical member {element} - {diagram_kind}\n"
        f"L = {length_m:.2f} m | Final status: {status} | Governing problem: {problem}",
        fontsize=11,
        fontweight="bold",
    )
    note = (
        "Exact station-by-station N/V/M data are plotted where supplied."
        if used_exact_data else
        "Only worst-result EC2 NEd/VEd/MEd ordinates were available; constant lines show governing demand ordinates, not inferred full-span diagrams."
    )
    fig.text(0.5, 0.012, note, ha="center", fontsize=7.1, style="italic")
    fig.tight_layout(rect=[0, 0.04, 1, 0.925])
    fig.savefig(filename, dpi=230, bbox_inches="tight")
    plt.close(fig)
    return filename


def save_failed_member_nmv_diagrams(
    diagnosis_df: Optional[pd.DataFrame],
    output_dir: str,
    *,
    nodes: Optional[dict] = None,
    elements: Optional[dict] = None,
    member_force_diagrams: Any = None,
    max_members: int = 12,
    filename_prefix: str = "07_failed_member",
) -> list[str]:
    """Save N/V/M plots for the most critical failed members."""
    if diagnosis_df is None or getattr(diagnosis_df, "empty", True):
        return []
    os.makedirs(output_dir, exist_ok=True)
    d = diagnosis_df.copy()
    critical_eta = pd.to_numeric(d.get("Eta_Final", pd.Series(np.nan, index=d.index)), errors="coerce")
    critical_eta = critical_eta.fillna(pd.to_numeric(d.get("Eta_After_Walls", pd.Series(np.nan, index=d.index)), errors="coerce"))
    d["_critical_eta"] = critical_eta
    if "Status_Final" in d.columns:
        d["_is_fail"] = d["Status_Final"].astype(str).str.upper().eq("FAIL")
    else:
        d["_is_fail"] = False
    d = d.sort_values(["_is_fail", "_critical_eta"], ascending=[False, False], na_position="last").head(max(1, int(max_members)))

    paths = []
    for _, row in d.iterrows():
        try:
            eid = int(row.get("Element"))
        except Exception:
            continue
        safe_prefix = str(filename_prefix or "07_failed_member").strip().replace(" ", "_")
        filename = os.path.join(output_dir, f"{safe_prefix}_{eid}_nmv_diagram.png")
        path = _plot_single_member_nmv(
            row,
            filename,
            nodes=nodes,
            elements=elements,
            member_force_diagrams=member_force_diagrams,
        )
        if path:
            paths.append(path)
    return paths


# ============================================================
# Figures: only professional/useful plots, no candidate-cost plot
# ============================================================


def save_story_drift_figure(df: Optional[pd.DataFrame], filename: str, drift_limit_ratio: Optional[float]) -> Optional[str]:
    if df is None or df.empty or "Story" not in df.columns:
        return None
    before = _first_col(df, ["Drift_Before_Ratio"])
    after = _first_col(df, ["Drift_After_Ratio"])
    if before is None and after is None:
        return None
    d = df.copy()
    d["Story"] = pd.to_numeric(d["Story"], errors="coerce")
    d = d.dropna(subset=["Story"]).sort_values("Story")
    if d.empty:
        return None
    plt.figure(figsize=(8.2, 5.2))
    if before is not None and pd.to_numeric(d[before], errors="coerce").notna().any():
        plt.plot(pd.to_numeric(d[before], errors="coerce"), d["Story"], marker="o", label="Before retrofit")
    if after is not None and pd.to_numeric(d[after], errors="coerce").notna().any():
        plt.plot(pd.to_numeric(d[after], errors="coerce"), d["Story"], marker="s", label="After selected shear walls")
    if drift_limit_ratio is not None:
        plt.axvline(float(drift_limit_ratio), linestyle="--", label="Drift limit")
    plt.xlabel("Story drift ratio [-]")
    plt.ylabel("Story")
    plt.title("Story Drift Envelope Before and After Retrofit")
    plt.grid(True, alpha=0.3)
    plt.legend()
    plt.tight_layout()
    plt.savefig(filename, dpi=220)
    plt.close()
    return filename


def save_member_eta_figure(df: Optional[pd.DataFrame], filename: str) -> Optional[str]:
    if df is None or df.empty or "Element" not in df.columns:
        return None
    cols = [c for c in ["Eta_Before", "Eta_After_Walls", "Eta_Final"] if c in df.columns]
    if not cols:
        return None
    d = df.copy()
    d["_critical"] = pd.to_numeric(d[cols[0]], errors="coerce")
    d = d.sort_values("_critical", ascending=False).head(30).sort_values("Element")
    x = np.arange(len(d))
    width = 0.8 / len(cols)
    plt.figure(figsize=(max(9, 0.35 * len(d)), 5.3))
    for i, c in enumerate(cols):
        plt.bar(x + (i - (len(cols) - 1) / 2) * width, pd.to_numeric(d[c], errors="coerce").fillna(0), width, label=c.replace("Eta_", ""))
    plt.axhline(1.0, linestyle="--", label="EC2 limit")
    plt.xticks(x, d["Element"].astype(str), rotation=90)
    plt.ylabel("Utilization ratio eta [-]")
    plt.xlabel("Element ID")
    plt.title("Critical Member Utilization Before and After Retrofit")
    plt.grid(True, axis="y", alpha=0.3)
    plt.legend()
    plt.tight_layout()
    plt.savefig(filename, dpi=220)
    plt.close()
    return filename



def _alternative_engineering_comment(row: pd.Series) -> str:
    """Create a concise engineering comment for one global retrofit alternative."""
    selected = str(row.get("Selected", "")).upper() == "YES"
    feasible = str(row.get("Complete_Feasible", "")).upper() == "YES"
    drift_ok = str(row.get("Drift_OK", "")).upper() == "YES"
    member_ok = str(row.get("Final_Member_OK", "")).upper() == "YES"
    final_fails = _safe_float(row.get("Final_Failed_Members", np.nan), np.nan)
    drift_util = _safe_float(row.get("Drift_Utilization", np.nan), np.nan)
    jacketing_required = str(row.get("Jacketing_Required", "")).upper() == "YES"
    limit_penalty = _safe_float(row.get("Penalty_Jacketing_Limit", 0.0), 0.0)

    if selected and feasible:
        return "Selected: lowest-cost complete feasible solution."
    if feasible:
        if jacketing_required:
            return "Feasible alternative; higher cost because member jacketing is still required."
        return "Feasible alternative; compare mainly by wall cost and drift margin."
    reasons = []
    if not drift_ok or (np.isfinite(drift_util) and drift_util > 1.0):
        reasons.append("drift not satisfied")
    if not member_ok or (np.isfinite(final_fails) and final_fails > 0):
        reasons.append("member failures remain")
    if limit_penalty > 0.0:
        reasons.append("jacketing/detailing limit reached")
    if not reasons:
        reasons.append("needs engineering review")
    prefix = "Selected (best available, NOT fully feasible): " if selected else "Not selected: "
    return prefix + "; ".join(reasons) + "."


def build_retrofit_alternatives_comparison(
    global_opt_df: Optional[pd.DataFrame],
    *,
    max_alternatives: int = 5,
) -> pd.DataFrame:
    """Return a compact Top-N comparison of complete retrofit alternatives.

    Each row is a complete strategy: selected shear-wall layout plus the
    concrete jacketing required by that layout. This is the main table for
    showing why the selected retrofit is preferable to competing alternatives.
    """
    if global_opt_df is None or getattr(global_opt_df, "empty", True):
        return pd.DataFrame()

    d_all = global_opt_df.copy()
    if "Strategy_Key" in d_all.columns:
        d_all = d_all.drop_duplicates(subset=["Strategy_Key"], keep="first")
    else:
        subset_cols = [c for c in ["Candidate_ID", "Wall_Description", "Description", "Jacketed_Members", "Total_Retrofit_Cost", "Drift_Utilization", "Final_Failed_Members"] if c in d_all.columns]
        if subset_cols:
            d_all = d_all.drop_duplicates(subset=subset_cols, keep="first")
    d = d_all.head(max(1, int(max_alternatives))).reset_index(drop=True)
    selected_cost = np.nan
    sel = d[d.get("Selected", "NO").astype(str).str.upper().eq("YES")] if "Selected" in d.columns else pd.DataFrame()
    if not sel.empty:
        selected_cost = _safe_float(sel.iloc[0].get("Total_Retrofit_Cost", np.nan), np.nan)
    elif "Total_Retrofit_Cost" in d.columns:
        selected_cost = _safe_float(d.iloc[0].get("Total_Retrofit_Cost", np.nan), np.nan)

    rows: list[dict[str, Any]] = []
    for i, row in d.iterrows():
        total = _safe_float(row.get("Total_Retrofit_Cost", np.nan), np.nan)
        delta = total - selected_cost if np.isfinite(total) and np.isfinite(selected_cost) else np.nan
        delta_pct = 100.0 * delta / selected_cost if np.isfinite(delta) and np.isfinite(selected_cost) and abs(selected_cost) > 1e-12 else np.nan
        rows.append({
            "Alternative": f"Alternative {i + 1}",
            "Selected": str(row.get("Selected", "")),
            "Candidate_ID": str(row.get("Candidate_ID", "")),
            "Wall_Description": str(row.get("Description", "")),
            "Strategy_Key": str(row.get("Strategy_Key", "")),
            "Complete_Feasible": str(row.get("Complete_Feasible", "")),
            "Wall_Cost": _safe_float(row.get("Wall_Cost", np.nan), np.nan),
            "Wall_Concrete_m3": _safe_float(row.get("Wall_Concrete_m3", row.get("Concrete_m3", np.nan)), np.nan),
            "Wall_Rebar_kg": _safe_float(row.get("Wall_Rebar_kg", row.get("Rebar_kg", np.nan)), np.nan),
            "Wall_Formwork_m2": _safe_float(row.get("Wall_Formwork_m2", row.get("Formwork_m2", np.nan)), np.nan),
            "Wall_Material_Cost": _safe_float(row.get("Wall_Material_Cost", np.nan), np.nan),
            "Wall_Labor_Cost": _safe_float(row.get("Wall_Labor_Cost", np.nan), np.nan),
            "Wall_Indirect_Cost": _safe_float(row.get("Wall_Indirect_Cost", np.nan), np.nan),
            "Wall_Foundation_Connection_Cost": _safe_float(row.get("Wall_Foundation_Connection_Cost", np.nan), np.nan),
            "Wall_Constructability_Cost": _safe_float(row.get("Wall_Constructability_Cost", np.nan), np.nan),
            "Wall_Site_Overhead_Cost": _safe_float(row.get("Wall_Site_Overhead_Cost", np.nan), np.nan),
            "Jacketing_Cost": _safe_float(row.get("Jacketing_Cost", np.nan), np.nan),
            "Jacketing_Concrete_m3": _safe_float(row.get("Jacketing_Concrete_m3", np.nan), np.nan),
            "Jacketing_Rebar_kg": _safe_float(row.get("Jacketing_Rebar_kg", np.nan), np.nan),
            "Jacketing_Formwork_m2": _safe_float(row.get("Jacketing_Formwork_m2", np.nan), np.nan),
            "Jacketing_Material_Cost": _safe_float(row.get("Jacketing_Material_Cost", np.nan), np.nan),
            "Jacketing_Labor_Cost": _safe_float(row.get("Jacketing_Labor_Cost", np.nan), np.nan),
            "Jacketing_Indirect_Cost": _safe_float(row.get("Jacketing_Indirect_Cost", np.nan), np.nan),
            "Jacketing_Site_Overhead_Cost": _safe_float(row.get("Jacketing_Site_Overhead_Cost", np.nan), np.nan),
            "Wall_Duration_Days": _safe_float(row.get("Wall_Duration_Days", np.nan), np.nan),
            "Jacketing_Duration_Days": _safe_float(row.get("Jacketing_Duration_Days", np.nan), np.nan),
            "Total_Duration_Days": _safe_float(row.get("Total_Duration_Days", np.nan), np.nan),
            "Total_Retrofit_Cost": total,
            "Cost_Difference_From_Selected": delta,
            "Cost_Difference_From_Selected_pct": delta_pct,
            "Drift_Utilization": _safe_float(row.get("Drift_Utilization", np.nan), np.nan),
            "Members_Failing_After_Walls": _safe_float(row.get("Member_Failures_After_Walls", np.nan), np.nan),
            "Jacketing_Required": str(row.get("Jacketing_Required", "")),
            "Final_Max_EC2_Utilization": _safe_float(row.get("Final_Max_EC2_Utilization", np.nan), np.nan),
            "Final_Failed_Members": _safe_float(row.get("Final_Failed_Members", np.nan), np.nan),
            "Constructability_Class": str(row.get("Constructability_Class", "")),
            "Global_Objective_Score": _safe_float(row.get("Global_Objective_Score", np.nan), np.nan),
            "Engineering_Comment": _alternative_engineering_comment(row),
        })

    return pd.DataFrame(rows)


def save_retrofit_alternative_cost_figure(df: Optional[pd.DataFrame], filename: str) -> Optional[str]:
    """Save a cost-comparison plot for the top retrofit alternatives."""
    if df is None or getattr(df, "empty", True):
        return None
    required = {"Alternative", "Wall_Cost", "Jacketing_Cost", "Total_Retrofit_Cost"}
    if not required.issubset(set(df.columns)):
        return None
    d = df.copy()
    labels = d["Alternative"].astype(str).tolist()
    x = np.arange(len(d))
    wall = pd.to_numeric(d["Wall_Cost"], errors="coerce").fillna(0.0)
    jack = pd.to_numeric(d["Jacketing_Cost"], errors="coerce").fillna(0.0)
    total = pd.to_numeric(d["Total_Retrofit_Cost"], errors="coerce").fillna(0.0)

    plt.figure(figsize=(8.5, 4.8))
    plt.bar(x, wall, label="Shear-wall cost")
    plt.bar(x, jack, bottom=wall, label="Concrete-jacketing cost")
    for i, value in enumerate(total):
        plt.text(i, value, f"{value:.0f}", ha="center", va="bottom", fontsize=8)
    plt.xticks(x, labels, rotation=20, ha="right")
    plt.ylabel("Estimated retrofit cost")
    plt.title("Top 5 Retrofit Alternatives - Cost Breakdown")
    plt.grid(True, axis="y", alpha=0.3)
    plt.legend()
    plt.tight_layout()
    plt.savefig(filename, dpi=230, bbox_inches="tight")
    plt.close()
    return filename


def save_retrofit_alternative_drift_cost_figure(df: Optional[pd.DataFrame], filename: str) -> Optional[str]:
    """Save a drift-utilization versus cost trade-off plot for the alternatives."""
    if df is None or getattr(df, "empty", True):
        return None
    if "Total_Retrofit_Cost" not in df.columns or "Drift_Utilization" not in df.columns:
        return None
    d = df.copy()
    cost = pd.to_numeric(d["Total_Retrofit_Cost"], errors="coerce")
    drift = pd.to_numeric(d["Drift_Utilization"], errors="coerce")
    mask = cost.notna() & drift.notna()
    if not mask.any():
        return None
    d = d.loc[mask].reset_index(drop=True)
    cost = cost.loc[mask].reset_index(drop=True)
    drift = drift.loc[mask].reset_index(drop=True)

    plt.figure(figsize=(8.2, 4.8))
    plt.scatter(cost, drift, s=70)
    for i, row in d.iterrows():
        plt.annotate(str(row.get("Alternative", f"A{i+1}")), (cost.iloc[i], drift.iloc[i]), xytext=(5, 4), textcoords="offset points", fontsize=8)
    plt.axhline(1.0, linestyle="--", label="Drift limit utilization = 1.0")
    plt.xlabel("Total retrofit cost")
    plt.ylabel("Drift utilization [-]")
    plt.title("Cost versus Drift Utilization for Top Retrofit Alternatives")
    plt.grid(True, alpha=0.3)
    plt.legend()
    plt.tight_layout()
    plt.savefig(filename, dpi=230, bbox_inches="tight")
    plt.close()
    return filename


def save_retrofit_alternative_failure_cost_figure(df: Optional[pd.DataFrame], filename: str) -> Optional[str]:
    """Save a final-failed-members versus cost trade-off plot."""
    if df is None or getattr(df, "empty", True):
        return None
    if "Total_Retrofit_Cost" not in df.columns or "Final_Failed_Members" not in df.columns:
        return None
    d = df.copy()
    cost = pd.to_numeric(d["Total_Retrofit_Cost"], errors="coerce")
    fails = pd.to_numeric(d["Final_Failed_Members"], errors="coerce")
    mask = cost.notna() & fails.notna()
    if not mask.any():
        return None
    d = d.loc[mask].reset_index(drop=True)
    cost = cost.loc[mask].reset_index(drop=True)
    fails = fails.loc[mask].reset_index(drop=True)

    plt.figure(figsize=(8.2, 4.6))
    plt.scatter(cost, fails, s=70)
    for i, row in d.iterrows():
        plt.annotate(str(row.get("Alternative", f"A{i+1}")), (cost.iloc[i], fails.iloc[i]), xytext=(5, 4), textcoords="offset points", fontsize=8)
    plt.axhline(0.0, linestyle="--", label="No final failed members")
    plt.xlabel("Total retrofit cost")
    plt.ylabel("Final failed members [count]")
    plt.title("Cost versus Remaining Member Failures")
    plt.grid(True, alpha=0.3)
    plt.legend()
    plt.tight_layout()
    plt.savefig(filename, dpi=230, bbox_inches="tight")
    plt.close()
    return filename




def _node_id_at_safe(geom: Any, floor: int, bay_index_zero_based: int) -> Optional[int]:
    """Return node id using framekit.geometry.node_id_at when available."""
    try:
        from framekit.geometry import node_id_at
        return int(node_id_at(geom, floor, bay_index_zero_based))
    except Exception:
        try:
            n_bays = int(getattr(geom, "num_bays"))
            return int(floor * (n_bays + 1) + bay_index_zero_based + 1)
        except Exception:
            return None


def _finalize_layout_axes(ax: Any, nodes: dict, title: str) -> None:
    ax.set_aspect("equal", adjustable="box")
    ax.set_title(title)
    ax.set_xlabel("X [m]")
    ax.set_ylabel("Y [m]")
    ax.grid(True, alpha=0.25)
    if nodes:
        xs = [float(x) for x, _ in nodes.values()]
        ys = [float(y) for _, y in nodes.values()]
        dx = max(xs) - min(xs) if max(xs) > min(xs) else 1.0
        dy = max(ys) - min(ys) if max(ys) > min(ys) else 1.0
        ax.set_xlim(min(xs) - 0.06 * dx, max(xs) + 0.06 * dx)
        ax.set_ylim(min(ys) - 0.08 * dy, max(ys) + 0.08 * dy)


def _draw_frame_layout(ax: Any, nodes: dict, elements: dict, *, member_ids: bool = True) -> None:
    for eid, (ni, nj, _etype) in elements.items():
        if ni not in nodes or nj not in nodes:
            continue
        x1, y1 = nodes[ni]
        x2, y2 = nodes[nj]
        ax.plot([x1, x2], [y1, y2], color="0.65", linewidth=1.4, zorder=1)
        if member_ids:
            ax.text((x1 + x2) / 2.0, (y1 + y2) / 2.0, str(eid), fontsize=6.5,
                    ha="center", va="center", color="0.25", zorder=3,
                    bbox=dict(facecolor="white", edgecolor="none", alpha=0.65, pad=0.5))
    for _nid, (x, y) in nodes.items():
        ax.plot(x, y, "k.", markersize=3.5, zorder=4)


def _draw_wall_panels(ax: Any, nodes: dict, geom: Any, walls: list, *, color: str, label: str, alpha: float) -> None:
    first = True
    for w in walls or []:
        try:
            ids = [
                _node_id_at_safe(geom, int(w.floor_bot), int(w.bay_left) - 1),
                _node_id_at_safe(geom, int(w.floor_bot), int(w.bay_right) - 1),
                _node_id_at_safe(geom, int(w.floor_top), int(w.bay_right) - 1),
                _node_id_at_safe(geom, int(w.floor_top), int(w.bay_left) - 1),
            ]
            if any(i is None or i not in nodes for i in ids):
                continue
            pts = [nodes[i] for i in ids]
            xs = [float(p[0]) for p in pts]
            ys = [float(p[1]) for p in pts]
            ax.fill(xs, ys, color=color, alpha=alpha, zorder=2, label=label if first else "_nolegend_")
            ax.plot(xs + [xs[0]], ys + [ys[0]], color=color, linewidth=2.0, zorder=3)
            ax.text(sum(xs) / 4.0, sum(ys) / 4.0, f"SW\nB{w.bay_left}\nS{w.floor_top}",
                    fontsize=7, ha="center", va="center", zorder=4,
                    bbox=dict(facecolor="white", edgecolor="none", alpha=0.72, pad=1.0))
            first = False
        except Exception:
            continue


def save_shear_wall_layout_figure(
    nodes: Optional[dict],
    elements: Optional[dict],
    geom: Any,
    filename: str,
    *,
    base_walls: Optional[list] = None,
    selected_candidate: Any = None,
) -> Optional[str]:
    """Save before/after structure layout with selected shear walls for Word/Excel reports."""
    if not nodes or not elements:
        return None
    new_walls = list(getattr(selected_candidate, "walls", []) or [])
    fig, axes = plt.subplots(1, 2, figsize=(12.5, 6.0), constrained_layout=True)
    _draw_frame_layout(axes[0], nodes, elements, member_ids=True)
    _draw_wall_panels(axes[0], nodes, geom, list(base_walls or []), color="0.35", label="Existing wall", alpha=0.20)
    _finalize_layout_axes(axes[0], nodes, "Before shear-wall retrofit")

    _draw_frame_layout(axes[1], nodes, elements, member_ids=True)
    _draw_wall_panels(axes[1], nodes, geom, list(base_walls or []), color="0.35", label="Existing wall", alpha=0.18)
    _draw_wall_panels(axes[1], nodes, geom, new_walls, color="tab:blue", label="Selected retrofit wall", alpha=0.30)
    title = "After selected shear-wall retrofit"
    desc = str(getattr(selected_candidate, "description", "") or "")
    if desc:
        title += "\n" + desc
    _finalize_layout_axes(axes[1], nodes, title)
    for ax in axes:
        handles, labels = ax.get_legend_handles_labels()
        if handles:
            ax.legend(fontsize=8)
    fig.savefig(filename, dpi=240, bbox_inches="tight")
    plt.close(fig)
    return filename


def save_jacketed_members_layout_figure(
    nodes: Optional[dict],
    elements: Optional[dict],
    jacketed_member_ids: Optional[list],
    filename: str,
) -> Optional[str]:
    """Save full-frame layout with every concrete-jacketed member highlighted in red."""
    if not nodes or not elements or not jacketed_member_ids:
        return None
    jacketed = {int(e) for e in jacketed_member_ids if str(e).strip()}
    if not jacketed:
        return None
    fig, ax = plt.subplots(figsize=(11.0, 6.5))
    ax.set_aspect("equal", adjustable="box")
    for eid, (ni, nj, etype) in elements.items():
        if ni not in nodes or nj not in nodes:
            continue
        x1, y1 = nodes[ni]
        x2, y2 = nodes[nj]
        is_j = int(eid) in jacketed
        ax.plot([x1, x2], [y1, y2], linewidth=5.2 if is_j else 1.4,
                color="tab:red" if is_j else "0.68", zorder=3 if is_j else 1)
        label = f"{eid}\nJACKET" if is_j else str(eid)
        ax.text((x1 + x2) / 2.0, (y1 + y2) / 2.0, label,
                ha="center", va="center", fontsize=7.5 if is_j else 6.5,
                color="tab:red" if is_j else "0.30",
                bbox=dict(facecolor="white", edgecolor="none", alpha=0.78, pad=0.6), zorder=4)
    for _nid, (x, y) in nodes.items():
        ax.plot(x, y, "k.", markersize=3.5, zorder=5)
    _finalize_layout_axes(ax, nodes, "Concrete jacketing layout - all strengthened members highlighted")
    fig.savefig(filename, dpi=240, bbox_inches="tight")
    plt.close(fig)
    return filename




def save_initial_structure_member_id_figure(
    nodes: Optional[dict],
    elements: Optional[dict],
    filename: str,
    *,
    title: str = "Initial structural model - member IDs",
) -> Optional[str]:
    """Save the undeformed initial frame layout with member IDs for the report."""
    if not nodes or not elements:
        return None
    fig, ax = plt.subplots(figsize=(11.0, 6.5))
    _draw_frame_layout(ax, nodes, elements, member_ids=True)
    _finalize_layout_axes(ax, nodes, title)
    fig.savefig(filename, dpi=240, bbox_inches="tight")
    plt.close(fig)
    return filename


def _fallback_utilization_color(eta: float):
    """Local color scale compatible with the EC2 utilization plot."""
    try:
        from framekit.ec2_checks import utilization_to_color
        return utilization_to_color(float(eta))
    except Exception:
        e = float(eta) if np.isfinite(float(eta)) else 0.0
        if e <= 0.70:
            return (0.75, 0.95, 0.75, 1.0)
        if e <= 1.00:
            return (1.00, 1.00, 0.70, 1.0)
        if e <= 1.20:
            return (1.00, 0.85, 0.60, 1.0)
        return (1.00, 0.70, 0.70, 1.0)


def _eta_map_from_worst_df(df_worst: Optional[pd.DataFrame]) -> dict[int, float]:
    """Build an Element -> eta map from the EC2 worst-utilization dataframe."""
    if df_worst is None or getattr(df_worst, "empty", True):
        return {}
    ele_col = _first_col(df_worst, ["Element", "ElementID", "Member", "Member ID"])
    eta_col = _first_col(df_worst, ["eta", "Static_EC2_Utilization", "Utilization", "Worst eta", "Max_Eta"])
    if ele_col is None or eta_col is None:
        return {}
    eta_map: dict[int, float] = {}
    for _, row in df_worst.iterrows():
        try:
            eid = int(pd.to_numeric(row[ele_col], errors="coerce"))
            eta = float(pd.to_numeric(row[eta_col], errors="coerce"))
        except Exception:
            continue
        if np.isfinite(eta):
            eta_map[eid] = eta
    return eta_map


def save_member_utilization_map_figure(
    nodes: Optional[dict],
    elements: Optional[dict],
    df_worst: Optional[pd.DataFrame],
    filename: str,
    *,
    title: str = "Baseline EC2 member utilization map before retrofit",
) -> Optional[str]:
    """Save a colored structural utilization map from a worst-utilization dataframe.

    This mirrors the interactive colored utilization plot, but creates a PNG so
    it can be inserted into the Word and Excel retrofit reports.
    """
    if not nodes or not elements:
        return None
    eta_map = _eta_map_from_worst_df(df_worst)
    if not eta_map:
        return None

    fig, ax = plt.subplots(figsize=(11.0, 6.5))
    ax.set_aspect("equal", adjustable="box")

    for eid, (ni, nj, _etype) in elements.items():
        if ni not in nodes or nj not in nodes:
            continue
        x1, y1 = nodes[ni]
        x2, y2 = nodes[nj]
        eta = eta_map.get(int(eid), np.nan)
        color = _fallback_utilization_color(eta) if np.isfinite(eta) else "0.78"
        lw = 4.6 if np.isfinite(eta) else 1.4
        ax.plot([x1, x2], [y1, y2], color=color, linewidth=lw, zorder=2)
        label = f"{eid}\n{eta:.2f}" if np.isfinite(eta) else str(eid)
        ax.text(
            (x1 + x2) / 2.0,
            (y1 + y2) / 2.0,
            label,
            ha="center",
            va="center",
            fontsize=7.2,
            color="0.15",
            bbox=dict(facecolor="white", edgecolor="none", alpha=0.78, pad=0.6),
            zorder=4,
        )

    for _nid, (x, y) in nodes.items():
        ax.plot(x, y, "k.", markersize=3.5, zorder=5)

    from matplotlib.lines import Line2D
    legend_items = [
        Line2D([0], [0], color=_fallback_utilization_color(0.50), lw=5, label="Low utilization"),
        Line2D([0], [0], color=_fallback_utilization_color(0.90), lw=5, label="Near capacity"),
        Line2D([0], [0], color=_fallback_utilization_color(1.10), lw=5, label="Above limit"),
    ]
    ax.legend(handles=legend_items, loc="best", fontsize=8, framealpha=0.9)
    _finalize_layout_axes(ax, nodes, title)
    fig.savefig(filename, dpi=240, bbox_inches="tight")
    plt.close(fig)
    return filename



def _jacketing_iteration_entry_to_df_and_label(entry: Any, fallback_iteration: int) -> tuple[Optional[pd.DataFrame], int, str]:
    """Normalize one saved jacketing-iteration EC2 table entry.

    Accepted forms:
    - pandas DataFrame
    - {"iteration": i, "df_worst": dataframe}
    - {"iteration": i, "df": dataframe}
    - (i, dataframe)
    """
    iteration = int(fallback_iteration)
    label = f"Jacketing iteration {iteration}"
    df = None

    if isinstance(entry, pd.DataFrame):
        df = entry
    elif isinstance(entry, dict):
        try:
            iteration = int(entry.get("iteration", fallback_iteration))
        except Exception:
            iteration = int(fallback_iteration)
        label = str(entry.get("label", f"Jacketing iteration {iteration}"))
        for key in ("df_worst", "df", "worst_df", "ec2_worst_df"):
            val = entry.get(key)
            if isinstance(val, pd.DataFrame):
                df = val
                break
    elif isinstance(entry, (list, tuple)) and len(entry) >= 2:
        try:
            iteration = int(entry[0])
        except Exception:
            iteration = int(fallback_iteration)
        label = f"Jacketing iteration {iteration}"
        if isinstance(entry[1], pd.DataFrame):
            df = entry[1]

    return df, iteration, label


def save_jacketing_iteration_utilization_maps(
    *,
    nodes: Optional[dict],
    elements: Optional[dict],
    iteration_worst_dfs: Optional[list[Any]],
    output_dir: str,
) -> list[str]:
    """Save one colored EC2 utilization map after each jacketing iteration."""
    if not nodes or not elements or not iteration_worst_dfs:
        return []

    os.makedirs(output_dir, exist_ok=True)
    paths: list[str] = []
    for fallback_i, entry in enumerate(iteration_worst_dfs, start=1):
        df, iteration, label = _jacketing_iteration_entry_to_df_and_label(entry, fallback_i)
        if df is None or getattr(df, "empty", True):
            continue
        filename = os.path.join(output_dir, f"05_jacketing_iteration_{iteration:02d}_member_utilization_map.png")
        path = save_member_utilization_map_figure(
            nodes,
            elements,
            df,
            filename,
            title=f"Member utilization after concrete jacketing - iteration {iteration}",
        )
        if path:
            paths.append(path)
    return paths

def _section_attr(sec: Any, names: list[str], default: float = 0.0) -> Any:
    for name in names:
        if hasattr(sec, name):
            return getattr(sec, name)
    return default


def _bar_positions_for_report(sec: Any, n: int, phi: float, y: float) -> list[tuple[float, float]]:
    b = float(_section_attr(sec, ["b_mm", "b"], 0.0))
    cover = float(_section_attr(sec, ["cover_mm", "cover"], 30.0))
    stirrup_phi = float(_section_attr(sec, ["stirrup_phi_mm", "stirrup_phi"], 8.0))
    x_left = cover + stirrup_phi + 0.5 * phi
    x_right = b - cover - stirrup_phi - 0.5 * phi
    if n <= 1:
        return [(0.5 * b, y)]
    return [(x_left + i * (x_right - x_left) / (n - 1), y) for i in range(n)]


def _draw_dimension_arrow(ax: Any, x1: float, y1: float, x2: float, y2: float, label: str, *, fontsize: float = 6.8) -> None:
    ax.annotate(
        "",
        xy=(x2, y2),
        xytext=(x1, y1),
        arrowprops=dict(arrowstyle="<->", lw=0.7, color="black", shrinkA=0, shrinkB=0),
    )
    xm = 0.5 * (x1 + x2)
    ym = 0.5 * (y1 + y2)
    if abs(y2 - y1) < 1e-9:
        ax.text(xm, ym + 5.0, label, ha="center", va="bottom", fontsize=fontsize)
    else:
        ax.text(xm + 5.0, ym, label, ha="left", va="center", rotation=90, fontsize=fontsize)


def _report_side_bar_positions_for_column(sec: Any, *, b: float, h: float, y_bot: float, y_top: float) -> list[tuple[float, float, float]]:
    """Return intermediate side-face bars for a column section.

    The top and bottom faces already include the corner bars through n_top and n_bot.
    n_left and n_right are therefore interpreted as intermediate side bars excluding
    the corner bars. This matches the section_library convention.
    """
    cover = float(_section_attr(sec, ["cover_mm", "cover"], 30.0))
    stirrup_phi = float(_section_attr(sec, ["stirrup_phi_mm", "stirrup_phi"], 8.0))
    phi_top = float(_section_attr(sec, ["phi_top_mm"], 0.0))
    phi_bot = float(_section_attr(sec, ["phi_bot_mm"], 0.0))
    phi_side = float(_section_attr(sec, ["phi_side_mm", "side_bar_d", "phi_side"], max(phi_top, phi_bot, 0.0)))
    n_left = int(_section_attr(sec, ["n_left", "left_bars", "n_side_left"], 0) or 0)
    n_right = int(_section_attr(sec, ["n_right", "right_bars", "n_side_right"], 0) or 0)

    # Side bars are placed on the same stirrup/bar centerline as the longitudinal bars.
    x_left = cover + stirrup_phi + 0.5 * phi_side
    x_right = b - cover - stirrup_phi - 0.5 * phi_side

    positions: list[tuple[float, float, float]] = []
    for side, n_side in (("left", n_left), ("right", n_right)):
        if n_side <= 0 or phi_side <= 0:
            continue
        x = x_left if side == "left" else x_right
        for i in range(n_side):
            # Divide the clear distance between bottom and top bar rows into n_side+1 spaces.
            y = y_bot + (i + 1) * (y_top - y_bot) / (n_side + 1)
            positions.append((x, y, phi_side))
    return positions


def _is_column_section_for_report(sec: Any, member_type: Optional[str] = None) -> bool:
    # If the caller knows the element type, trust it. This prevents beam
    # sections that carry generic side-bar attributes from being drawn as
    # column sections in the report.
    if member_type is not None and str(member_type).strip():
        return "column" in str(member_type).strip().lower()
    role = str(_section_attr(sec, ["section_role", "role", "member_type"], "")).lower()
    if "column" in role:
        return True
    return any(hasattr(sec, name) for name in ["n_left", "n_right", "phi_side_mm", "left_bars", "right_bars"])


def _draw_rc_section_for_report(ax: Any, sec: Any, subtitle: str, *, original_sec: Any = None, member_type: Optional[str] = None) -> None:
    """Draw a report-style RC section without coordinate axes.

    Existing concrete is white. Added jacketing concrete is solid gray.
    Dimensions are shown using conventional dimension arrows. Column sections
    are drawn with longitudinal bars around all faces when side-bar attributes
    are available.
    """
    from matplotlib.patches import Rectangle, Circle

    b = float(_section_attr(sec, ["b_mm", "b"], 0.0))
    h = float(_section_attr(sec, ["h_mm", "h"], 0.0))
    cover = float(_section_attr(sec, ["cover_mm", "cover"], 30.0))
    stirrup_phi = float(_section_attr(sec, ["stirrup_phi_mm", "stirrup_phi"], 8.0))
    n_top = int(_section_attr(sec, ["n_top", "top_bars"], 0) or 0)
    n_bot = int(_section_attr(sec, ["n_bot", "bottom_bars"], 0) or 0)
    phi_top = float(_section_attr(sec, ["phi_top_mm", "top_bar_d"], 0.0))
    phi_bot = float(_section_attr(sec, ["phi_bot_mm", "bottom_bar_d"], 0.0))
    phi_side = float(_section_attr(sec, ["phi_side_mm", "side_bar_d", "phi_side"], max(phi_top, phi_bot, 0.0)))
    n_left = int(_section_attr(sec, ["n_left", "left_bars", "n_side_left"], 0) or 0)
    n_right = int(_section_attr(sec, ["n_right", "right_bars", "n_side_right"], 0) or 0)
    spacing = float(_section_attr(sec, ["stirrup_spacing_mm", "stirrup_spacing"], 0.0))
    legs = int(_section_attr(sec, ["stirrup_legs"], 0) or 0)
    is_column = _is_column_section_for_report(sec, member_type=member_type)

    ax.set_aspect("equal", adjustable="box")
    ax.axis("off")

    if original_sec is None:
        ax.add_patch(Rectangle((0, 0), b, h, facecolor="white", edgecolor="black", linewidth=1.4, zorder=1))
        extra = f"cover = {cover:.0f} mm"
    else:
        ob = float(_section_attr(original_sec, ["b_mm", "b"], 0.0))
        oh = float(_section_attr(original_sec, ["h_mm", "h"], 0.0))
        ox = 0.5 * (b - ob)
        oy = 0.5 * (h - oh)
        t = max(0.5 * (b - ob), 0.5 * (h - oh), 0.0)
        ax.add_patch(Rectangle((0, 0), b, h, facecolor="0.78", edgecolor="black", linewidth=1.4, zorder=1))
        ax.add_patch(Rectangle((ox, oy), ob, oh, facecolor="white", edgecolor="black", linewidth=1.1, zorder=2))
        ax.text(ox + 0.5 * ob, oy + oh + 0.025 * h, "existing", ha="center", va="bottom", fontsize=6.0, color="0.25")
        extra = f"jacket t = {t:.0f} mm"

    s_off = cover + 0.5 * stirrup_phi
    ax.add_patch(Rectangle((s_off, s_off), max(b - 2.0 * s_off, 1.0), max(h - 2.0 * s_off, 1.0),
                           fill=False, linestyle="--", linewidth=0.95, edgecolor="0.35", zorder=3))

    y_top = h - cover - stirrup_phi - 0.5 * phi_top
    y_bot = cover + stirrup_phi + 0.5 * phi_bot

    # Top and bottom rows include the corner bars. Side bars are intermediate bars.
    for x, y in _bar_positions_for_report(sec, n_top, phi_top, y_top):
        if phi_top > 0:
            ax.add_patch(Circle((x, y), max(0.5 * phi_top, 1.2), fill=True, color="black", zorder=4))
    for x, y in _bar_positions_for_report(sec, n_bot, phi_bot, y_bot):
        if phi_bot > 0:
            ax.add_patch(Circle((x, y), max(0.5 * phi_bot, 1.2), fill=True, color="black", zorder=4))

    if is_column:
        for x, y, phi in _report_side_bar_positions_for_column(sec, b=b, h=h, y_bot=y_bot, y_top=y_top):
            ax.add_patch(Circle((x, y), max(0.5 * phi, 1.2), fill=True, color="black", zorder=4))

    m = max(b, h) * 0.28
    y_dim = h + 0.12 * m
    x_dim = b + 0.13 * m
    _draw_dimension_arrow(ax, 0, y_dim, b, y_dim, f"{b:.0f} mm")
    _draw_dimension_arrow(ax, x_dim, 0, x_dim, h, f"{h:.0f} mm")
    ax.plot([0, 0], [h, y_dim + 0.04 * m], color="black", lw=0.55)
    ax.plot([b, b], [h, y_dim + 0.04 * m], color="black", lw=0.55)
    ax.plot([b, x_dim + 0.04 * m], [0, 0], color="black", lw=0.55)
    ax.plot([b, x_dim + 0.04 * m], [h, h], color="black", lw=0.55)

    if is_column:
        rebar_note = (
            f"Top {n_top}Ø{phi_top:.0f} | Bot {n_bot}Ø{phi_bot:.0f}\n"
            f"Side L/R {n_left}+{n_right}Ø{phi_side:.0f}"
        )
    else:
        rebar_note = f"Top {n_top}Ø{phi_top:.0f} | Bot {n_bot}Ø{phi_bot:.0f}"

    ax.set_title(subtitle, fontsize=8, fontweight="bold", pad=3)
    ax.text(
        0.5 * b,
        -0.20 * m,
        f"b = {b:.0f} mm   h = {h:.0f} mm\n{rebar_note}\nStirrups Ø{stirrup_phi:.0f}@{spacing:.0f}, legs={legs}\n{extra}",
        ha="center",
        va="top",
        fontsize=6.2,
    )
    ax.set_xlim(-0.22 * m, b + 0.35 * m)
    ax.set_ylim(-0.44 * m, h + 0.38 * m)


def save_jacketed_section_comparison_figures(
    *,
    elements: Optional[dict],
    column_sec: Any,
    beam_sec: Any,
    jacketed_element_sections: Optional[dict],
    jacketed_member_ids: Optional[list],
    output_dir: str,
    members_per_figure: int = 3,
    filename_prefix: str = "05_jacketed_section_comparisons",
    figure_title: str = "Concrete Jacketing Section Comparisons",
) -> list[str]:
    """Save before/after concrete-jacketing section comparison figures.

    This makes the section comparison plots that you see interactively available
    inside the retrofit_report folder and inserts them into Word/Excel through
    the existing report figure pipeline.
    """
    if not elements or jacketed_element_sections is None or not jacketed_member_ids:
        return []

    os.makedirs(output_dir, exist_ok=True)

    valid_ids: list[int] = []
    for eid in jacketed_member_ids:
        try:
            eid_int = int(eid)
        except Exception:
            continue
        if eid_int in elements and eid_int in jacketed_element_sections:
            valid_ids.append(eid_int)

    valid_ids = sorted(set(valid_ids))
    if not valid_ids:
        return []

    paths: list[str] = []
    members_per_figure = max(1, int(members_per_figure or 3))

    for page, start_idx in enumerate(range(0, len(valid_ids), members_per_figure), start=1):
        chunk = valid_ids[start_idx:start_idx + members_per_figure]
        nrows = len(chunk)
        fig_height = max(4.2, 3.4 * nrows)
        fig, axes = plt.subplots(nrows, 2, figsize=(8.6, fig_height), squeeze=False)

        for row, eid in enumerate(chunk):
            try:
                _ni, _nj, etype = elements[eid]
            except Exception:
                continue
            etype_txt = str(etype).lower()
            base_sec = column_sec if etype_txt == "column" else beam_sec
            new_sec = jacketed_element_sections[eid]

            _draw_rc_section_for_report(
                axes[row][0],
                base_sec,
                f"Element {eid} before ({etype})",
                member_type=etype_txt,
            )
            _draw_rc_section_for_report(
                axes[row][1],
                new_sec,
                f"Element {eid} after jacketing",
                original_sec=base_sec,
                member_type=etype_txt,
            )

        fig.suptitle(str(figure_title or "Concrete Jacketing Section Comparisons"), fontsize=12, fontweight="bold")
        fig.text(
            0.5,
            0.015,
            "White = existing section     |     Solid gray = added concrete jacket",
            ha="center",
            fontsize=8,
        )
        fig.tight_layout(rect=[0, 0.035, 1, 0.955])
        safe_prefix = str(filename_prefix or "05_jacketed_section_comparisons").strip().replace(" ", "_")
        fig_path = os.path.join(output_dir, f"{safe_prefix}_{page:02d}.png")
        fig.savefig(fig_path, dpi=240, bbox_inches="tight")
        plt.close(fig)
        paths.append(fig_path)

    return paths


# ============================================================
# Detailed Top-5 alternative packages
# ============================================================

def _parse_int_list(value: Any) -> list[int]:
    """Parse comma/semicolon separated member IDs into sorted integers."""
    if value is None:
        return []
    if isinstance(value, (list, tuple, set)):
        vals = value
    else:
        txt = str(value).replace(";", ",").replace("|", ",")
        vals = [x.strip() for x in txt.split(",")]
    out: list[int] = []
    for item in vals:
        try:
            if str(item).strip():
                out.append(int(float(str(item).strip())))
        except Exception:
            continue
    return sorted(set(out))


def _short_text(value: Any, max_len: int = 220) -> str:
    txt = str(value or "")
    return txt if len(txt) <= max_len else txt[: max_len - 3] + "..."


def _top_unique_global_rows(global_opt_df: Optional[pd.DataFrame], max_alternatives: int = 5) -> list[pd.Series]:
    if global_opt_df is None or getattr(global_opt_df, "empty", True):
        return []
    d = global_opt_df.copy()
    if "Strategy_Key" in d.columns:
        d = d.drop_duplicates(subset=["Strategy_Key"], keep="first")
    return [row for _, row in d.head(max(1, int(max_alternatives))).iterrows()]


def _alt_value(row: pd.Series, names: list[str], default: Any = np.nan) -> Any:
    for name in names:
        try:
            val = row.get(name, None)
        except Exception:
            val = None
        if val is not None:
            try:
                if pd.isna(val):
                    continue
            except Exception:
                pass
            return val
    return default


def _build_alternative_summary_table(row: pd.Series, alt_no: int) -> pd.DataFrame:
    """Compact summary for one complete retrofit alternative."""
    walls = list(getattr(row.get("_candidate", None), "walls", []) or [])
    wall_panels = []
    for w in walls:
        try:
            wall_panels.append(f"Bay {int(w.bay_left)}: floor {int(w.floor_bot)}-{int(w.floor_top)}")
        except Exception:
            continue
    jacketed = str(row.get("Jacketed_Members", "") or "")
    data = [
        ("Alternative", f"Alternative {alt_no}"),
        ("Selected", str(row.get("Selected", ""))),
        ("Candidate ID", str(row.get("Candidate_ID", ""))),
        ("Description", _short_text(row.get("Description", ""), 500)),
        ("Complete feasible", str(row.get("Complete_Feasible", ""))),
        ("Suggested shear-wall panels", "; ".join(wall_panels) if wall_panels else "No new shear walls"),
        ("Wall thickness [mm]", _safe_float(row.get("Wall_Thickness_mm", np.nan))),
        ("Number of wall panels", _safe_float(row.get("Num_Wall_Panels", np.nan))),
        ("Number of wall lines", _safe_float(row.get("Num_Wall_Lines", row.get("Wall_Lines", np.nan)))),
        ("Jacketed members", jacketed if jacketed else "No concrete jacketing"),
        ("Drift utilization", _safe_float(row.get("Drift_Utilization", np.nan))),
        ("Members failing after walls", _safe_float(row.get("Member_Failures_After_Walls", np.nan))),
        ("Final max EC2 utilization", _safe_float(row.get("Final_Max_EC2_Utilization", np.nan))),
        ("Final failed members", _safe_float(row.get("Final_Failed_Members", np.nan))),
        ("Constructability class", str(row.get("Constructability_Class", ""))),
        ("Engineering comment", _alternative_engineering_comment(row)),
    ]
    return pd.DataFrame(data, columns=["Item", "Value"])


def _cost_param(cost_parameters: Optional[dict], key: str, default: float = 0.0) -> float:
    try:
        return float((cost_parameters or {}).get(key, default))
    except Exception:
        return float(default)


def _build_alternative_cost_formula_table(row: pd.Series, cost_parameters: Optional[dict]) -> pd.DataFrame:
    """Show the actual cost formula used for one alternative."""
    u = dict(cost_parameters or {})
    wc = _safe_float(_alt_value(row, ["Wall_Concrete_m3", "Concrete_m3"], 0.0), 0.0)
    wr = _safe_float(_alt_value(row, ["Wall_Rebar_kg", "Rebar_kg"], 0.0), 0.0)
    wf = _safe_float(_alt_value(row, ["Wall_Formwork_m2", "Formwork_m2"], 0.0), 0.0)
    wl = _safe_float(_alt_value(row, ["Num_Wall_Lines", "Wall_Lines"], 0.0), 0.0)
    wd = _safe_float(row.get("Wall_Duration_Days", 0.0), 0.0)
    jc = _safe_float(row.get("Jacketing_Concrete_m3", 0.0), 0.0)
    jr = _safe_float(row.get("Jacketing_Rebar_kg", 0.0), 0.0)
    jf = _safe_float(row.get("Jacketing_Formwork_m2", 0.0), 0.0)
    jd = _safe_float(row.get("Jacketing_Duration_Days", 0.0), 0.0)

    wall_conc_unit = _cost_param(u, "wall_concrete_per_m3")
    wall_rebar_unit = _cost_param(u, "wall_rebar_per_kg")
    wall_form_unit = _cost_param(u, "wall_formwork_per_m2")
    jacket_conc_unit = _cost_param(u, "jacket_concrete_per_m3")
    jacket_rebar_unit = _cost_param(u, "jacket_rebar_per_kg")
    jacket_form_unit = _cost_param(u, "jacket_formwork_per_m2")
    labor_factor = _cost_param(u, "labor_factor")
    indirect_factor = _cost_param(u, "indirect_factor")
    foundation_unit = _cost_param(u, "foundation_cost_per_wall_line")
    constructability_unit = _cost_param(u, "constructability_cost_per_wall_line")
    overhead_unit = _cost_param(u, "site_overhead_per_day")

    wall_material = wc * wall_conc_unit + wr * wall_rebar_unit + wf * wall_form_unit
    wall_labor = labor_factor * wall_material
    wall_indirect = indirect_factor * wall_material
    wall_foundation = wl * foundation_unit
    wall_constructability = wl * constructability_unit
    wall_overhead = wd * overhead_unit
    wall_total_formula = wall_material + wall_labor + wall_indirect + wall_foundation + wall_constructability + wall_overhead
    wall_total_reported = _safe_float(row.get("Wall_Cost", wall_total_formula), wall_total_formula)

    jacket_material = jc * jacket_conc_unit + jr * jacket_rebar_unit + jf * jacket_form_unit
    jacket_labor = labor_factor * jacket_material
    jacket_indirect = indirect_factor * jacket_material
    jacket_overhead = jd * overhead_unit
    jacket_total_formula = jacket_material + jacket_labor + jacket_indirect + jacket_overhead
    jacket_total_reported = _safe_float(row.get("Jacketing_Cost", jacket_total_formula), jacket_total_formula)
    total_reported = _safe_float(row.get("Total_Retrofit_Cost", wall_total_reported + jacket_total_reported), wall_total_reported + jacket_total_reported)

    rows = [
        {
            "Stage": "Shear walls", "Component": "Concrete material", "Quantity": wc, "Unit": "m3",
            "Unit_Cost": wall_conc_unit, "Formula": f"{wc:.3f} m3 x {wall_conc_unit:.2f}", "Cost": wc * wall_conc_unit,
        },
        {
            "Stage": "Shear walls", "Component": "Reinforcement material", "Quantity": wr, "Unit": "kg",
            "Unit_Cost": wall_rebar_unit, "Formula": f"{wr:.3f} kg x {wall_rebar_unit:.2f}", "Cost": wr * wall_rebar_unit,
        },
        {
            "Stage": "Shear walls", "Component": "Formwork", "Quantity": wf, "Unit": "m2",
            "Unit_Cost": wall_form_unit, "Formula": f"{wf:.3f} m2 x {wall_form_unit:.2f}", "Cost": wf * wall_form_unit,
        },
        {"Stage": "Shear walls", "Component": "Material subtotal", "Quantity": np.nan, "Unit": "", "Unit_Cost": np.nan, "Formula": "concrete + rebar + formwork", "Cost": wall_material},
        {"Stage": "Shear walls", "Component": "Labor", "Quantity": labor_factor, "Unit": "ratio", "Unit_Cost": np.nan, "Formula": f"{labor_factor:.3f} x wall material subtotal", "Cost": wall_labor},
        {"Stage": "Shear walls", "Component": "Indirect", "Quantity": indirect_factor, "Unit": "ratio", "Unit_Cost": np.nan, "Formula": f"{indirect_factor:.3f} x wall material subtotal", "Cost": wall_indirect},
        {"Stage": "Shear walls", "Component": "Foundation/connection allowance", "Quantity": wl, "Unit": "wall lines", "Unit_Cost": foundation_unit, "Formula": f"{wl:.0f} wall lines x {foundation_unit:.2f}", "Cost": wall_foundation},
        {"Stage": "Shear walls", "Component": "Constructability allowance", "Quantity": wl, "Unit": "wall lines", "Unit_Cost": constructability_unit, "Formula": f"{wl:.0f} wall lines x {constructability_unit:.2f}", "Cost": wall_constructability},
        {"Stage": "Shear walls", "Component": "Site overhead", "Quantity": wd, "Unit": "days", "Unit_Cost": overhead_unit, "Formula": f"{wd:.2f} days x {overhead_unit:.2f}", "Cost": wall_overhead},
        {"Stage": "Shear walls", "Component": "Wall cost used by optimizer", "Quantity": np.nan, "Unit": "", "Unit_Cost": np.nan, "Formula": "sum of wall components", "Cost": wall_total_reported},
        {
            "Stage": "Concrete jacketing", "Component": "Concrete material", "Quantity": jc, "Unit": "m3",
            "Unit_Cost": jacket_conc_unit, "Formula": f"{jc:.3f} m3 x {jacket_conc_unit:.2f}", "Cost": jc * jacket_conc_unit,
        },
        {
            "Stage": "Concrete jacketing", "Component": "Reinforcement material", "Quantity": jr, "Unit": "kg",
            "Unit_Cost": jacket_rebar_unit, "Formula": f"{jr:.3f} kg x {jacket_rebar_unit:.2f}", "Cost": jr * jacket_rebar_unit,
        },
        {
            "Stage": "Concrete jacketing", "Component": "Formwork", "Quantity": jf, "Unit": "m2",
            "Unit_Cost": jacket_form_unit, "Formula": f"{jf:.3f} m2 x {jacket_form_unit:.2f}", "Cost": jf * jacket_form_unit,
        },
        {"Stage": "Concrete jacketing", "Component": "Material subtotal", "Quantity": np.nan, "Unit": "", "Unit_Cost": np.nan, "Formula": "concrete + rebar + formwork", "Cost": jacket_material},
        {"Stage": "Concrete jacketing", "Component": "Labor", "Quantity": labor_factor, "Unit": "ratio", "Unit_Cost": np.nan, "Formula": f"{labor_factor:.3f} x jacket material subtotal", "Cost": jacket_labor},
        {"Stage": "Concrete jacketing", "Component": "Indirect", "Quantity": indirect_factor, "Unit": "ratio", "Unit_Cost": np.nan, "Formula": f"{indirect_factor:.3f} x jacket material subtotal", "Cost": jacket_indirect},
        {"Stage": "Concrete jacketing", "Component": "Site overhead", "Quantity": jd, "Unit": "days", "Unit_Cost": overhead_unit, "Formula": f"{jd:.2f} days x {overhead_unit:.2f}", "Cost": jacket_overhead},
        {"Stage": "Concrete jacketing", "Component": "Jacketing cost used by optimizer", "Quantity": np.nan, "Unit": "", "Unit_Cost": np.nan, "Formula": "sum of jacketing components", "Cost": jacket_total_reported},
        {"Stage": "Total", "Component": "Total retrofit cost", "Quantity": np.nan, "Unit": "", "Unit_Cost": np.nan, "Formula": "wall cost + jacketing cost", "Cost": total_reported},
    ]
    return pd.DataFrame(rows)


def _build_alternative_member_report(row: pd.Series) -> pd.DataFrame:
    rpt = row.get("_member_report", None)
    out = _normalize_jacketing_member_report(rpt)
    if out.empty:
        ids = _parse_int_list(row.get("Jacketed_Members", ""))
        if ids:
            out = pd.DataFrame({"Element": ids, "Note": "Jacketed member in this alternative"})
    return out


def build_alternative_detail_packages(
    global_opt_df: Optional[pd.DataFrame],
    *,
    cost_parameters: Optional[dict],
    output_dir: str,
    nodes: Optional[dict],
    elements: Optional[dict],
    geom: Any,
    base_walls: Optional[list],
    column_sec: Any,
    beam_sec: Any,
    df_worst_before: Optional[pd.DataFrame],
    member_force_diagrams: Any = None,
    max_alternatives: int = 5,
) -> list[dict[str, Any]]:
    """Create tables and figures for a detailed section for each Top-N alternative."""
    packages: list[dict[str, Any]] = []
    rows = _top_unique_global_rows(global_opt_df, max_alternatives=max_alternatives)
    for alt_no, row in enumerate(rows, start=1):
        selected = str(row.get("Selected", "")).upper() == "YES"
        candidate = row.get("_candidate", None)
        jacketed_ids = _parse_int_list(row.get("Jacketed_Members", ""))
        element_sections = row.get("_element_sections", None)
        df_after = row.get("_df_worst_after_walls", None)
        df_final = row.get("_df_worst_final", df_after)
        member_fail_after = row.get("_member_fail_after_walls", None)
        member_fail_final = row.get("_member_fail_final", None)
        member_report = _build_alternative_member_report(row)

        alt_dir = os.path.join(output_dir, "top5_alternative_details")
        os.makedirs(alt_dir, exist_ok=True)
        prefix = f"alt_{alt_no:02d}"
        figures: list[str] = []

        sw_fig = save_shear_wall_layout_figure(
            nodes, elements, geom, os.path.join(alt_dir, f"{prefix}_shear_wall_layout.png"),
            base_walls=base_walls, selected_candidate=candidate,
        )
        if sw_fig:
            figures.append(sw_fig)
        jk_fig = save_jacketed_members_layout_figure(
            nodes, elements, jacketed_ids, os.path.join(alt_dir, f"{prefix}_jacketed_members_layout.png")
        )
        if jk_fig:
            figures.append(jk_fig)
        sec_figs = save_jacketed_section_comparison_figures(
            elements=elements,
            column_sec=column_sec,
            beam_sec=beam_sec,
            jacketed_element_sections=element_sections,
            jacketed_member_ids=jacketed_ids,
            output_dir=alt_dir,
            members_per_figure=3,
            filename_prefix=f"{prefix}_jacketed_section_comparisons",
            figure_title=f"Alternative {alt_no} - Jacketed Section Comparisons",
        )
        figures.extend(sec_figs)

        member_compare = build_member_utilization_comparison(
            df_worst_before,
            df_after,
            df_final,
            jacketing_member_report=row.get("_member_report", None),
            global_opt_df=global_opt_df,
        )
        diagnosis = build_failed_member_diagnosis_table(
            member_compare_df=member_compare,
            df_worst_after_walls=df_after,
            df_worst_final=df_final,
            member_fail_after_walls=member_fail_after,
            member_fail_final=member_fail_final,
        )
        # For alternative-by-alternative comparison, limit N/V/M diagrams to the
        # most critical few members so the Word file stays readable. The final
        # selected decision section still receives the full selected-member set.
        nmv_figs = save_failed_member_nmv_diagrams(
            diagnosis,
            alt_dir,
            nodes=nodes,
            elements=elements,
            member_force_diagrams=member_force_diagrams if selected else None,
            max_members=3,
            filename_prefix=f"{prefix}_critical_member",
        )
        figures.extend(nmv_figs)

        packages.append({
            "alternative_no": alt_no,
            "title": f"Alternative {alt_no}" + (" - SELECTED" if selected else ""),
            "selected": selected,
            "summary_df": _build_alternative_summary_table(row, alt_no),
            "cost_formula_df": _build_alternative_cost_formula_table(row, cost_parameters),
            "member_report_df": member_report,
            "diagnosis_df": diagnosis,
            "figures": figures,
        })
    return packages


# ============================================================
# Excel and Word export
# ============================================================


def _xlsx_col(idx: int) -> str:
    idx += 1
    out = ""
    while idx:
        idx, rem = divmod(idx - 1, 26)
        out = chr(65 + rem) + out
    return out


def _write_sheet(writer: pd.ExcelWriter, df: Optional[pd.DataFrame], name: str) -> None:
    if df is None or getattr(df, "empty", True):
        return
    sheet = name[:31]
    df.to_excel(writer, sheet_name=sheet, index=False)
    wb = writer.book
    ws = writer.sheets[sheet]
    header = wb.add_format({"bold": True, "text_wrap": True, "align": "center", "valign": "vcenter", "bg_color": "#D9EAF7", "border": 1})
    cell = wb.add_format({"border": 1})
    num = wb.add_format({"border": 1, "num_format": "0.000000"})
    for j, col in enumerate(df.columns):
        ws.write(0, j, col, header)
        width = min(max(len(str(col)) + 2, 14), 45)
        ws.set_column(j, j, width, num if pd.api.types.is_numeric_dtype(df[col]) else cell)
    ws.freeze_panes(1, 0)
    for j, col in enumerate(df.columns):
        c = str(col).lower()
        rng = f"{_xlsx_col(j)}2:{_xlsx_col(j)}{len(df) + 1}"
        if "status" in c or "feasible" in c:
            ws.conditional_format(rng, {"type": "text", "criteria": "containing", "value": "OK", "format": wb.add_format({"bg_color": "#C6EFCE", "font_color": "#006100"})})
            ws.conditional_format(rng, {"type": "text", "criteria": "containing", "value": "FAIL", "format": wb.add_format({"bg_color": "#FFC7CE", "font_color": "#9C0006"})})
        if "eta" in c or "drift" in c or "ratio" in c:
            ws.conditional_format(rng, {"type": "3_color_scale"})


def _fmt(value: Any) -> str:
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    if isinstance(value, (float, np.floating)):
        if abs(float(value)) >= 1e6 or (0 < abs(float(value)) < 1e-4):
            return f"{float(value):.3e}"
        if abs(float(value)) >= 1000:
            return f"{float(value):,.2f}"
        return f"{float(value):.6g}"
    return str(value)


def _add_table(doc: Any, df: Optional[pd.DataFrame], title: str, max_rows: int = 35, max_cols: int = 9) -> None:
    if df is None or getattr(df, "empty", True):
        return
    from docx.shared import Pt
    from docx.enum.table import WD_TABLE_ALIGNMENT
    doc.add_heading(title, level=2)
    show = df.head(max_rows).copy()
    preferred = [
        "Metric", "Item", "Value", "Stage", "Component", "Quantity", "Unit", "Unit_Cost", "Formula", "Cost",
        "Story", "Governing_Case_Before", "Drift_Before_Ratio",
        "Governing_Case_After", "Drift_After_Ratio", "Drift_Reduction_pct",
        "Element", "Type", "Eta_Before", "Eta_Before_Shear", "Eta_Before_Flexure_Axial", "Problem_Before",
        "Eta_After_Walls", "Eta_Final", "Eta_Final_Shear", "Eta_Final_Flexure_Axial", "Problem_After",
        "Eta_Reduction_pct", "Chosen_Jacket_Thickness_mm", "Chosen_Stirrup_Spacing_mm",
        "Chosen_Longitudinal_Rebar_Increase_pct", "Final_Status",
        "Iteration", "Jacket_Thickness_mm", "Max_Allowed_Jacket_Thickness_By_Member_mm",
        "Max_Jacketed_Section_Dimension_Ratio", "Members_At_Jacket_Size_Limit",
        "Alternative_Required_Members", "Alternative_Recommendation",
        "Total_Cost", "Candidate_ID", "Description", "Feasible",
        "Alternative", "Selected", "Candidate_ID", "Wall_Description", "Strategy_Key", "Complete_Feasible",
        "Wall_Cost", "Jacketing_Cost", "Total_Retrofit_Cost", "Cost_Difference_From_Selected",
        "Cost_Difference_From_Selected_pct", "Drift_Utilization", "Members_Failing_After_Walls",
        "Jacketing_Required", "Final_Max_EC2_Utilization", "Final_Failed_Members",
        "Engineering_Comment", "Objective_Rank", "Objective_Total_Score", "Objective_Cost_Score", "Objective_Drift_Score",
        "Objective_Constructability_Score", "Objective_Wall_Quantity_Score", "Constructability_Class",
    ]
    cols = [c for c in preferred if c in show.columns] + [c for c in show.columns if c not in preferred]
    show = show[cols[:max_cols]]
    table = doc.add_table(rows=1, cols=len(show.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    fs = 7.0 if len(show.columns) > 7 else 8.5
    for j, col in enumerate(show.columns):
        cell = table.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(str(col).replace("_", " "))
        r.bold = True
        r.font.size = Pt(fs)
    for _, row in show.iterrows():
        cells = table.add_row().cells
        for j, col in enumerate(show.columns):
            text = _fmt(row[col])
            if len(text) > 80:
                text = text[:77] + "..."
            cells[j].text = ""
            r = cells[j].paragraphs[0].add_run(text)
            r.font.size = Pt(fs)
    if len(df) > max_rows or len(df.columns) > max_cols:
        p = doc.add_paragraph("Note: table truncated in Word report; complete data are provided in the Excel file.")
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(8)



def _add_picture_fit(doc: Any, image_path: str, *, max_width_in: float = 6.3, max_height_in: float = 5.2, align_center: bool = True) -> bool:
    """Insert an image into a Word document while keeping it inside page margins.

    The helper preserves the image aspect ratio and limits both width and height.
    This prevents tall section-comparison figures from becoming too large in
    the report. If Pillow is unavailable, it falls back to a conservative width.
    """
    if not image_path or not os.path.exists(image_path):
        return False
    try:
        from docx.shared import Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        width_px = height_px = None
        try:
            from PIL import Image
            with Image.open(image_path) as img:
                width_px, height_px = img.size
        except Exception:
            pass

        paragraph = doc.add_paragraph()
        if align_center:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()

        if width_px and height_px and width_px > 0 and height_px > 0:
            aspect = float(width_px) / float(height_px)
            width_in = min(float(max_width_in), float(max_height_in) * aspect)
            height_in = width_in / aspect
            if height_in > float(max_height_in):
                height_in = float(max_height_in)
                width_in = height_in * aspect
            run.add_picture(image_path, width=Inches(width_in), height=Inches(height_in))
        else:
            run.add_picture(image_path, width=Inches(min(float(max_width_in), 6.0)))
        return True
    except Exception:
        return False


def _resolve_artiste_logo_path(output_dir: str, logo_path: Optional[str] = None, artiste_logo_path: Optional[str] = None) -> Optional[str]:
    """Find the ArtIStE cover image near the Polito logo, report folder, or working directory."""
    names = [
        "artiste_logo_ridotto-768x512.jpg",
        "artiste_logo_ridotto.jpg",
        "ArtIStE.jpg",
        "Artiste.jpg",
        "artiste.jpg",
        "artiste_logo.png",
    ]
    candidates: list[str] = []
    if artiste_logo_path:
        candidates.append(str(artiste_logo_path))
    if logo_path:
        logo_dir = os.path.dirname(os.path.abspath(str(logo_path)))
        candidates += [os.path.join(logo_dir, name) for name in names]
    candidates += [os.path.join(os.getcwd(), name) for name in names]
    candidates += [os.path.join(output_dir, name) for name in names]
    for p in candidates:
        if p and os.path.exists(p):
            return p
    return None

def _create_word_report(
    output_dir: str,
    executive_df: pd.DataFrame,
    drift_compare_df: pd.DataFrame,
    member_compare_df: pd.DataFrame,
    cost_parameters_df: Optional[pd.DataFrame],
    sw_export_df: Optional[pd.DataFrame],
    global_export_df: Optional[pd.DataFrame],
    top_alternatives_df: Optional[pd.DataFrame],
    alternative_detail_packages: Optional[list[dict[str, Any]]],
    optimization_dataset_df: Optional[pd.DataFrame],
    jacketing_trials_df: Optional[pd.DataFrame],
    jacketing_member_report: Optional[pd.DataFrame],
    failed_member_diagnosis_df: Optional[pd.DataFrame],
    member_fail_after_walls: Optional[pd.DataFrame],
    member_fail_final: Optional[pd.DataFrame],
    figures: list[str],
    logo_path: Optional[str],
    artiste_logo_path: Optional[str],
    report_title: str,
    student_name: str,
    student_role: str,
    professor_name: str,
    research_topic: str,
) -> Optional[str]:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except Exception as exc:
        print(f"Word report not created because python-docx is unavailable: {exc}")
        return None

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

    if logo_path and os.path.exists(logo_path):
        _add_picture_fit(doc, logo_path, max_width_in=5.7, max_height_in=1.35)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("\n" + report_title)
    r.bold = True
    r.font.size = Pt(20)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(research_topic)
    r.italic = True
    r.font.size = Pt(13)
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run(f"Student: {student_name}\n").bold = True
    info.add_run(f"Role: {student_role}\n")
    info.add_run(f"Supervisor: Prof. {professor_name}\n")
    info.add_run(f"Date: {date.today().isoformat()}")

    # Add the ArtIStE visual at the bottom of the cover page when available.
    # It is resolved automatically from the same folder as the Polito logo,
    # the report output folder, or the current working directory.
    resolved_artiste_logo = _resolve_artiste_logo_path(output_dir, logo_path, artiste_logo_path)
    if resolved_artiste_logo:
        spacer = doc.add_paragraph()
        spacer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacer.add_run("\n")
        _add_picture_fit(doc, resolved_artiste_logo, max_width_in=6.45, max_height_in=2.65)

    doc.add_page_break()

    doc.add_heading("1. Executive Summary", level=1)
    _add_table(doc, executive_df, "Key Retrofit Metrics", max_rows=60, max_cols=2)

    doc.add_heading("2. Initial Structural Model and Baseline Member Utilization", level=1)
    for fig in figures:
        base = os.path.basename(fig).lower() if fig else ""
        if fig and os.path.exists(fig) and "initial_structure_member_ids" in base:
            doc.add_heading("Initial Structural Model with Member IDs", level=2)
            _add_picture_fit(doc, fig, max_width_in=6.35, max_height_in=5.1)
        if fig and os.path.exists(fig) and "initial_member_utilization_map" in base:
            doc.add_heading("Baseline Colored EC2 Member Utilization Map", level=2)
            _add_picture_fit(doc, fig, max_width_in=6.35, max_height_in=5.1)

    doc.add_heading("3. Baseline and Retrofitted Drift Performance", level=1)
    for fig in figures:
        if fig and os.path.exists(fig) and os.path.basename(fig).lower().startswith("01_story_drift"):
            _add_picture_fit(doc, fig, max_width_in=6.25, max_height_in=4.95)
    _add_table(doc, drift_compare_df, "Story-by-Story Drift Envelope", max_rows=60)

    doc.add_heading("4. Global Retrofit Alternatives and Cost Comparison", level=1)
    _add_table(doc, cost_parameters_df, "Cost Assumptions Used in Optimization", max_rows=20, max_cols=3)
    doc.add_paragraph(
        "The global optimizer compares complete retrofit alternatives: each alternative contains a shear-wall layout plus the concrete jacketing required after that layout. "
        "The selected strategy is the lowest-cost complete feasible solution satisfying drift and EC2 member-safety requirements; "
        "if no alternative satisfies all requirements, the best available alternative is selected and flagged as NOT fully feasible."
    )
    _add_table(doc, top_alternatives_df, "Top 5 Complete Retrofit Alternatives", max_rows=5, max_cols=14)

    if alternative_detail_packages:
        doc.add_heading("Detailed Review of the Top 5 Alternatives", level=2)
        doc.add_paragraph(
            "Each subsection below is a complete retrofit strategy. The report shows the suggested shear-wall layout, "
            "the concrete-jacketed members, the after-retrofit section drawings where available, and the explicit cost formula used by the optimizer."
        )
        for pkg in alternative_detail_packages:
            doc.add_heading(str(pkg.get("title", "Alternative")), level=3)
            _add_table(doc, pkg.get("summary_df"), "Alternative Layout and Performance Summary", max_rows=25, max_cols=2)
            _add_table(doc, pkg.get("cost_formula_df"), "Alternative Cost Calculation Formula", max_rows=30, max_cols=7)
            for fig in pkg.get("figures", []) or []:
                base = os.path.basename(fig).lower() if fig else ""
                if fig and os.path.exists(fig) and "shear_wall_layout" in base:
                    doc.add_heading("Suggested Shear-Wall Location", level=4)
                    _add_picture_fit(doc, fig, max_width_in=6.35, max_height_in=4.8)
                elif fig and os.path.exists(fig) and "jacketed_members_layout" in base:
                    doc.add_heading("Suggested Concrete-Jacketing Locations", level=4)
                    _add_picture_fit(doc, fig, max_width_in=6.35, max_height_in=4.8)
                elif fig and os.path.exists(fig) and "jacketed_section_comparisons" in base:
                    doc.add_heading("After-Retrofit Section Details", level=4)
                    _add_picture_fit(doc, fig, max_width_in=6.15, max_height_in=4.85)
                elif fig and os.path.exists(fig) and "critical_member" in base and "nmv_diagram" in base:
                    doc.add_heading("Critical Member N/V/M Demand Diagram", level=4)
                    _add_picture_fit(doc, fig, max_width_in=6.15, max_height_in=4.85)
            _add_table(doc, pkg.get("member_report_df"), "Alternative Jacketing Member Decisions", max_rows=25, max_cols=12)
            _add_table(doc, pkg.get("diagnosis_df"), "Alternative Critical Member Diagnosis", max_rows=20, max_cols=12)

    for fig in figures:
        base = os.path.basename(fig).lower() if fig else ""
        if fig and os.path.exists(fig) and "alternative_cost_breakdown" in base:
            doc.add_heading("Cost Breakdown of the Top Alternatives", level=2)
            _add_picture_fit(doc, fig, max_width_in=6.35, max_height_in=4.9)
        if fig and os.path.exists(fig) and "alternative_drift_cost" in base:
            doc.add_heading("Cost versus Drift Utilization", level=2)
            _add_picture_fit(doc, fig, max_width_in=6.35, max_height_in=4.9)
        if fig and os.path.exists(fig) and "alternative_failure_cost" in base:
            doc.add_heading("Cost versus Remaining Member Failures", level=2)
            _add_picture_fit(doc, fig, max_width_in=6.35, max_height_in=4.9)
    _add_table(doc, global_export_df, "Full Global Retrofit Strategy Ranking", max_rows=20, max_cols=12)
    _add_table(doc, optimization_dataset_df, "Optimization Dataset Preview", max_rows=12, max_cols=10)
    for fig in figures:
        if fig and os.path.exists(fig) and os.path.basename(fig).lower().startswith("03_shear_wall_layout"):
            doc.add_heading("Structure Layout Before and After Selected Shear Walls", level=2)
            _add_picture_fit(doc, fig, max_width_in=6.35, max_height_in=5.1)
    _add_table(doc, sw_export_df, "Shear Wall Candidate Summary", max_rows=25, max_cols=14)

    doc.add_heading("5. Concrete Jacketing Optimization", level=1)
    for fig in figures:
        base = os.path.basename(fig).lower() if fig else ""
        if fig and os.path.exists(fig) and base.startswith("04_jacketed_members_layout"):
            doc.add_heading("Concrete Jacketing Layout - Strengthened Members", level=2)
            _add_picture_fit(doc, fig, max_width_in=6.35, max_height_in=5.1)

    utilization_sequence = []
    for fig in figures:
        base = os.path.basename(fig).lower() if fig else ""
        if fig and os.path.exists(fig) and (
            "member_utilization_map_after_shear_walls" in base
            or "jacketing_iteration_" in base
        ):
            utilization_sequence.append(fig)
    if utilization_sequence:
        doc.add_heading("Member Utilization Evolution During Retrofitting", level=2)
        for fig in utilization_sequence:
            base = os.path.basename(fig).lower()
            if "after_shear_walls" in base:
                doc.add_heading("After Selected Shear Walls", level=3)
            elif "jacketing_iteration_" in base:
                label = base.replace("05_", "", 1).replace("_member_utilization_map.png", "").replace("_", " ").title()
                doc.add_heading(label, level=3)
            _add_picture_fit(doc, fig, max_width_in=6.35, max_height_in=5.1)

    for fig in figures:
        base = os.path.basename(fig).lower() if fig else ""
        if fig and os.path.exists(fig) and base.startswith("05_jacketed_section_comparisons"):
            doc.add_heading("Before/After Section Comparison for Jacketed Members", level=2)
            _add_picture_fit(doc, fig, max_width_in=6.15, max_height_in=4.85)
    _add_table(doc, jacketing_trials_df, "Jacketing Iteration History", max_rows=40)
    _add_table(doc, jacketing_member_report, "Per-Member Jacketing Diagnosis and Decisions", max_rows=45, max_cols=18)
    _add_table(doc, member_fail_after_walls, "Members Failing After Shear Walls", max_rows=30)
    _add_table(doc, member_fail_final, "Members Failing After Final Jacketing", max_rows=30)

    doc.add_heading("6. Failed-Member Diagnosis and N/M/V Demand Diagrams", level=1)
    doc.add_paragraph(
        "This section focuses on members that failed after the shear-wall stage or remained critical after jacketing. "
        "The diagnosis compares shear utilization with flexure/axial utilization and links the governing mechanism to the selected retrofit action. "
        "N/V/M diagrams are generated from exact station-by-station OpenSees member-force data when available. "
        "If such data are not supplied, the report clearly falls back to EC2 worst-result demand ordinates only."
    )
    _add_table(doc, failed_member_diagnosis_df, "Critical Member Diagnosis and Recommended Action", max_rows=40, max_cols=12)
    for fig in figures:
        base = os.path.basename(fig).lower() if fig else ""
        if fig and os.path.exists(fig) and "failed_member_" in base and "_nmv_diagram" in base:
            title = base.replace("07_", "", 1).replace("_nmv_diagram.png", "").replace("_", " ").title()
            doc.add_heading(title, level=2)
            _add_picture_fit(doc, fig, max_width_in=6.15, max_height_in=4.85)

    doc.add_heading("7. Final Member Verification", level=1)
    for fig in figures:
        base = os.path.basename(fig).lower() if fig else ""
        if fig and os.path.exists(fig) and "final_member_utilization_map" in base:
            doc.add_heading("Final Colored Member Utilization Map", level=2)
            _add_picture_fit(doc, fig, max_width_in=6.35, max_height_in=5.1)
    for fig in figures:
        base = os.path.basename(fig).lower() if fig else ""
        if fig and os.path.exists(fig) and base.startswith("02_member_utilization"):
            _add_picture_fit(doc, fig, max_width_in=6.25, max_height_in=4.95)
    _add_table(doc, member_compare_df, "Member Utilization Comparison", max_rows=45)

    doc.add_heading("8. Final Selected Retrofit Decision", level=1)
    selected_pkg = None
    if alternative_detail_packages:
        selected_pkg = next((p for p in alternative_detail_packages if p.get("selected")), alternative_detail_packages[0])
    if selected_pkg is not None:
        doc.add_paragraph(
            "This section collects the final selected strategy in one place, including the cost calculation, wall/jacketing locations, "
            "after-retrofit section drawings and critical N/V/M demand diagrams."
        )
        _add_table(doc, selected_pkg.get("summary_df"), "Selected Strategy Summary", max_rows=25, max_cols=2)
        _add_table(doc, selected_pkg.get("cost_formula_df"), "Selected Strategy Cost Calculation", max_rows=30, max_cols=7)
        for fig in selected_pkg.get("figures", []) or []:
            base = os.path.basename(fig).lower() if fig else ""
            if fig and os.path.exists(fig) and (
                "shear_wall_layout" in base or "jacketed_members_layout" in base or "jacketed_section_comparisons" in base or "nmv_diagram" in base
            ):
                _add_picture_fit(doc, fig, max_width_in=6.2, max_height_in=4.85)
    for fig in figures:
        base = os.path.basename(fig).lower() if fig else ""
        if fig and os.path.exists(fig) and "failed_member_" in base and "_nmv_diagram" in base:
            # Full selected critical-member diagrams from the normal final decision pipeline.
            _add_picture_fit(doc, fig, max_width_in=6.15, max_height_in=4.85)

    doc.add_heading("9. Engineering Conclusion", level=1)
    doc.add_paragraph(
        "The report uses the original static drift dataframe as the source of truth for the baseline condition, "
        "then compares it with the selected shear-wall candidate and the final member verification after concrete jacketing. "
        "Detailed tabular results are provided in the accompanying Excel workbook. The Top 5 alternative sections now show the cost formula and detailing information used to compare the candidate strategies."
    )

    path = os.path.join(output_dir, "retrofit_optimization_report.docx")
    doc.save(path)
    return path




def build_report_optimization_dataset(global_opt_df: Optional[pd.DataFrame], structure_id: str = "current_structure") -> pd.DataFrame:
    """Create a CSV/Excel-friendly dataset of all evaluated retrofit alternatives."""
    if global_opt_df is None or getattr(global_opt_df, "empty", True):
        return pd.DataFrame()
    df = global_opt_df.copy()
    df = df.drop(columns=[c for c in df.columns if str(c).startswith("_")], errors="ignore")
    preferred = [
        "Global_Rank", "Alternative_Number", "Selected", "Complete_Feasible",
        "Candidate_ID", "Description", "Strategy_Key",
        "Wall_Layout_Mode", "Wall_Segments", "Wall_Thickness_mm", "Num_Wall_Panels", "Num_Wall_Lines",
        "Constructability_Class", "Max_Drift_Ratio", "Drift_Limit_Ratio", "Drift_Utilization", "Drift_OK",
        "Member_Failures_After_Walls", "Jacketing_Required", "Jacketed_Members",
        "Final_Max_EC2_Utilization", "Final_Failed_Members", "Final_Member_OK",
        "Wall_Concrete_m3", "Wall_Rebar_kg", "Wall_Formwork_m2",
        "Wall_Material_Cost", "Wall_Labor_Cost", "Wall_Indirect_Cost",
        "Wall_Foundation_Connection_Cost", "Wall_Constructability_Cost", "Wall_Site_Overhead_Cost",
        "Jacketing_Concrete_m3", "Jacketing_Rebar_kg", "Jacketing_Formwork_m2",
        "Jacketing_Material_Cost", "Jacketing_Labor_Cost", "Jacketing_Indirect_Cost",
        "Jacketing_Site_Overhead_Cost", "Jacketing_Num_Members",
        "Wall_Cost", "Jacketing_Cost", "Total_Retrofit_Cost",
        "Wall_Duration_Days", "Jacketing_Duration_Days", "Total_Duration_Days",
        "Penalty_Drift", "Penalty_Member", "Penalty_Constructability", "Penalty_Jacketing_Limit", "Global_Objective_Score",
        "Jacketing_Convergence", "Alternative_Required_Members", "Selection_Explanation", "Cost_Assumptions",
    ]
    ordered = [c for c in preferred if c in df.columns]
    remaining = [c for c in df.columns if c not in ordered]
    df = df[ordered + remaining].copy()
    df.insert(0, "Dataset_Record_ID", [f"ALT_{i:04d}" for i in range(1, len(df) + 1)])
    df.insert(1, "Structure_ID", str(structure_id or "current_structure"))
    return df

def cost_parameters_to_dataframe(cost_parameters: Optional[dict]) -> pd.DataFrame:
    """Convert retrofit cost assumptions to a report-friendly table."""
    if not cost_parameters:
        return pd.DataFrame()
    labels = [
        ("wall_concrete_per_m3", "Wall concrete cost", "EUR/m3"),
        ("wall_rebar_per_kg", "Wall reinforcement cost", "EUR/kg"),
        ("wall_formwork_per_m2", "Wall formwork cost", "EUR/m2"),
        ("jacket_concrete_per_m3", "Jacketing concrete cost", "EUR/m3"),
        ("jacket_rebar_per_kg", "Jacketing reinforcement cost", "EUR/kg"),
        ("jacket_formwork_per_m2", "Jacketing formwork cost", "EUR/m2"),
        ("labor_factor", "Labor factor", "%"),
        ("indirect_factor", "Indirect factor", "%"),
        ("foundation_cost_per_wall_line", "Foundation/connection allowance", "EUR/wall line"),
        ("constructability_cost_per_wall_line", "Constructability allowance", "EUR/wall line"),
        ("site_overhead_per_day", "Site overhead", "EUR/day"),
        ("wall_productivity_m3_per_day", "Wall productivity", "m3/day"),
        ("jacketing_productivity_m3_per_day", "Jacketing volume productivity", "m3/day"),
        ("jacketing_productivity_members_per_day", "Jacketing member productivity", "members/day"),
    ]
    rows = []
    for key, label, unit in labels:
        if key not in cost_parameters:
            continue
        val = cost_parameters.get(key)
        if key in {"labor_factor", "indirect_factor"}:
            val = 100.0 * float(val)
        rows.append({"Parameter": label, "Value": val, "Unit": unit})
    return pd.DataFrame(rows)

def export_retrofit_optimization_report(
    *,
    df_static_drifts_before: Optional[pd.DataFrame] = None,
    sw_opt_df: Optional[pd.DataFrame] = None,
    global_opt_df: Optional[pd.DataFrame] = None,
    optimization_dataset_df: Optional[pd.DataFrame] = None,
    df_worst_before: Optional[pd.DataFrame] = None,
    df_worst_after_walls: Optional[pd.DataFrame] = None,
    df_worst_final: Optional[pd.DataFrame] = None,
    member_fail_after_walls: Optional[pd.DataFrame] = None,
    member_fail_final: Optional[pd.DataFrame] = None,
    jacketing_trials_df: Optional[pd.DataFrame] = None,
    jacketing_member_report: Optional[pd.DataFrame] = None,
    retrofit_cost_parameters: Optional[dict] = None,
    jacket_cost: Any = None,
    nodes: Optional[dict] = None,
    elements: Optional[dict] = None,
    geom: Any = None,
    base_walls: Optional[list] = None,
    selected_shear_wall_candidate: Any = None,
    jacketed_member_ids: Optional[list] = None,
    jacketed_element_sections: Optional[dict] = None,
    jacketing_iteration_worst_dfs: Optional[list[Any]] = None,
    column_sec: Any = None,
    beam_sec: Any = None,
    member_force_diagrams: Any = None,
    drift_limit_ratio: Optional[float] = None,
    output_dir: str = "retrofit_report",
    logo_path: Optional[str] = None,
    artiste_logo_path: Optional[str] = None,
    student_name: str = "Arman Hajiha",
    student_role: str = "Research assistant at Politecnico di Torino",
    professor_name: str = "Giuseppe Carlo Marano",
    report_title: str = "Seismic Assessment and Retrofit Optimization Report",
    research_topic: str = "AI and Structural Engineering",
) -> dict[str, Any]:
    os.makedirs(output_dir, exist_ok=True)

    drift_compare_df = build_story_drift_comparison(df_static_drifts_before, sw_opt_df)
    final_df = df_worst_final if df_worst_final is not None and not getattr(df_worst_final, "empty", True) else df_worst_after_walls
    member_compare_df = build_member_utilization_comparison(
        df_worst_before,
        df_worst_after_walls,
        final_df,
        jacketing_member_report=jacketing_member_report,
        global_opt_df=global_opt_df,
    )
    jacketing_member_report_df = _normalize_jacketing_member_report(jacketing_member_report)
    failed_member_diagnosis_df = build_failed_member_diagnosis_table(
        member_compare_df=member_compare_df,
        df_worst_after_walls=df_worst_after_walls,
        df_worst_final=final_df,
        member_fail_after_walls=member_fail_after_walls,
        member_fail_final=member_fail_final,
    )

    executive_df = build_executive_summary(
        drift_compare_df=drift_compare_df,
        drift_limit_ratio=drift_limit_ratio,
        sw_opt_df=sw_opt_df,
        df_worst_before=df_worst_before,
        df_worst_after_walls=df_worst_after_walls,
        df_worst_final=final_df,
        jacketing_trials_df=jacketing_trials_df,
        jacket_cost=jacket_cost,
        jacketing_member_report=jacketing_member_report,
        global_opt_df=global_opt_df,
    )

    sw_export_df = sw_opt_df.drop(columns=[c for c in sw_opt_df.columns if str(c).startswith("_")] + ["Selected_Drifts"], errors="ignore") if sw_opt_df is not None else None
    global_export_df = global_opt_df.drop(columns=[c for c in global_opt_df.columns if str(c).startswith("_")], errors="ignore") if global_opt_df is not None else None
    top_alternatives_df = build_retrofit_alternatives_comparison(global_opt_df, max_alternatives=5)
    if optimization_dataset_df is None or getattr(optimization_dataset_df, "empty", True):
        optimization_dataset_df = build_report_optimization_dataset(global_opt_df, structure_id="current_structure")
    cost_parameters_df = cost_parameters_to_dataframe(retrofit_cost_parameters)

    alternative_detail_packages = build_alternative_detail_packages(
        global_opt_df,
        cost_parameters=retrofit_cost_parameters,
        output_dir=output_dir,
        nodes=nodes,
        elements=elements,
        geom=geom,
        base_walls=base_walls,
        column_sec=column_sec,
        beam_sec=beam_sec,
        df_worst_before=df_worst_before,
        member_force_diagrams=member_force_diagrams,
        max_alternatives=5,
    )
    alternative_detail_figs = [
        fig
        for pkg in alternative_detail_packages
        for fig in (pkg.get("figures", []) or [])
        if fig
    ]

    initial_structure_fig = save_initial_structure_member_id_figure(
        nodes, elements, os.path.join(output_dir, "00_initial_structure_member_ids.png")
    )
    initial_utilization_fig = save_member_utilization_map_figure(
        nodes, elements, df_worst_before, os.path.join(output_dir, "00_initial_member_utilization_map_before_retrofit.png"),
        title="Initial EC2 member utilization map before retrofit",
    )
    after_walls_utilization_fig = save_member_utilization_map_figure(
        nodes, elements, df_worst_after_walls, os.path.join(output_dir, "03_member_utilization_map_after_shear_walls.png"),
        title="Member utilization after selected shear walls",
    )
    jacketing_iteration_utilization_figs = save_jacketing_iteration_utilization_maps(
        nodes=nodes,
        elements=elements,
        iteration_worst_dfs=jacketing_iteration_worst_dfs,
        output_dir=output_dir,
    )
    final_utilization_fig = save_member_utilization_map_figure(
        nodes, elements, final_df, os.path.join(output_dir, "06_final_member_utilization_map_after_retrofit.png"),
        title="Final member utilization after shear walls + concrete jacketing",
    )
    drift_fig = save_story_drift_figure(drift_compare_df, os.path.join(output_dir, "01_story_drift_envelope_before_after.png"), drift_limit_ratio)
    member_fig = save_member_eta_figure(member_compare_df, os.path.join(output_dir, "02_member_utilization_before_after.png"))
    alternative_cost_fig = save_retrofit_alternative_cost_figure(
        top_alternatives_df,
        os.path.join(output_dir, "02a_alternative_cost_breakdown_top5.png"),
    )
    alternative_drift_cost_fig = save_retrofit_alternative_drift_cost_figure(
        top_alternatives_df,
        os.path.join(output_dir, "02b_alternative_drift_cost_tradeoff_top5.png"),
    )
    alternative_failure_cost_fig = save_retrofit_alternative_failure_cost_figure(
        top_alternatives_df,
        os.path.join(output_dir, "02c_alternative_failure_cost_tradeoff_top5.png"),
    )
    failed_member_nmv_figs = save_failed_member_nmv_diagrams(
        failed_member_diagnosis_df,
        output_dir,
        nodes=nodes,
        elements=elements,
        member_force_diagrams=member_force_diagrams,
        max_members=12,
    )

    if selected_shear_wall_candidate is None and sw_opt_df is not None and not getattr(sw_opt_df, "empty", True):
        try:
            selected_shear_wall_candidate = sw_opt_df.iloc[0].get("_candidate", None)
        except Exception:
            selected_shear_wall_candidate = None

    sw_layout_fig = save_shear_wall_layout_figure(
        nodes, elements, geom, os.path.join(output_dir, "03_shear_wall_layout_before_after.png"),
        base_walls=base_walls, selected_candidate=selected_shear_wall_candidate,
    )
    jacket_layout_fig = save_jacketed_members_layout_figure(
        nodes, elements, jacketed_member_ids, os.path.join(output_dir, "04_jacketed_members_layout.png")
    )
    jacket_section_figs = save_jacketed_section_comparison_figures(
        elements=elements,
        column_sec=column_sec,
        beam_sec=beam_sec,
        jacketed_element_sections=jacketed_element_sections,
        jacketed_member_ids=jacketed_member_ids,
        output_dir=output_dir,
    )
    figures = [p for p in [initial_structure_fig, initial_utilization_fig, after_walls_utilization_fig, drift_fig, member_fig, alternative_cost_fig, alternative_drift_cost_fig, alternative_failure_cost_fig, sw_layout_fig, jacket_layout_fig, final_utilization_fig] if p]
    figures += jacketing_iteration_utilization_figs + jacket_section_figs + failed_member_nmv_figs + alternative_detail_figs

    xlsx_path = os.path.join(output_dir, "retrofit_optimization_report.xlsx")
    try:
        import xlsxwriter  # noqa: F401
        _xlsx_engine = "xlsxwriter"
    except Exception:
        _xlsx_engine = "openpyxl"
    with pd.ExcelWriter(xlsx_path, engine=_xlsx_engine) as writer:
        _write_sheet(writer, executive_df, "Executive_Summary")
        _write_sheet(writer, drift_compare_df, "Story_Drift_Comparison")
        _write_sheet(writer, member_compare_df, "Member_Utilization")
        _write_sheet(writer, cost_parameters_df, "Cost_Assumptions")
        _write_sheet(writer, top_alternatives_df, "Top_5_Retrofit_Alternatives")
        for pkg in alternative_detail_packages:
            alt_no = int(pkg.get("alternative_no", 0) or 0)
            _write_sheet(writer, pkg.get("summary_df"), f"Alt_{alt_no}_Summary")
            _write_sheet(writer, pkg.get("cost_formula_df"), f"Alt_{alt_no}_Cost_Formula")
            _write_sheet(writer, pkg.get("member_report_df"), f"Alt_{alt_no}_Jacketing")
            _write_sheet(writer, pkg.get("diagnosis_df"), f"Alt_{alt_no}_Diagnosis")
        _write_sheet(writer, global_export_df, "Global_Retrofit_Optimization")
        _write_sheet(writer, optimization_dataset_df, "Optimization_Dataset")
        _write_sheet(writer, sw_export_df, "Shear_Wall_Optimization")
        _write_sheet(writer, jacketing_trials_df, "Jacketing_Trials")
        _write_sheet(writer, jacketing_member_report_df, "Jacketing_Member_Report")
        _write_sheet(writer, failed_member_diagnosis_df, "Failed_Member_Diagnosis")
        _write_sheet(writer, df_worst_before, "EC2_Before")
        _write_sheet(writer, df_worst_after_walls, "EC2_After_Walls")
        _write_sheet(writer, final_df, "EC2_Final")
        _write_sheet(writer, member_fail_after_walls, "Fail_After_Walls")
        _write_sheet(writer, member_fail_final, "Fail_Final")
        if _xlsx_engine == "xlsxwriter":
            wb = writer.book
            ws = wb.add_worksheet("Figures")
            writer.sheets["Figures"] = ws
            row = 0
            for path in figures:
                if path and os.path.exists(path):
                    title = os.path.splitext(os.path.basename(path))[0].replace("_", " ").title()
                    ws.write(row, 0, title, wb.add_format({"bold": True}))
                    try:
                        from PIL import Image as _PILImage
                        with _PILImage.open(path) as _im:
                            w_px, h_px = _im.size
                    except Exception:
                        w_px, h_px = 1000, 750
                    sc = min(0.72, 900.0 / max(w_px, 1))
                    ws.insert_image(row + 1, 0, path, {"x_scale": sc, "y_scale": sc})
                    row += int(h_px * sc / 20.0) + 3

    docx_path = _create_word_report(
        output_dir=output_dir,
        executive_df=executive_df,
        drift_compare_df=drift_compare_df,
        member_compare_df=member_compare_df,
        cost_parameters_df=cost_parameters_df,
        sw_export_df=sw_export_df,
        global_export_df=global_export_df,
        top_alternatives_df=top_alternatives_df,
        alternative_detail_packages=alternative_detail_packages,
        optimization_dataset_df=optimization_dataset_df,
        jacketing_trials_df=jacketing_trials_df,
        jacketing_member_report=jacketing_member_report_df,
        failed_member_diagnosis_df=failed_member_diagnosis_df,
        member_fail_after_walls=member_fail_after_walls,
        member_fail_final=member_fail_final,
        figures=figures,
        logo_path=logo_path,
        artiste_logo_path=artiste_logo_path,
        report_title=report_title,
        student_name=student_name,
        student_role=student_role,
        professor_name=professor_name,
        research_topic=research_topic,
    )

    return {
        "output_dir": output_dir,
        "xlsx": xlsx_path,
        "docx": docx_path,
        "executive_summary": executive_df,
        "drift_comparison": drift_compare_df,
        "member_comparison": member_compare_df,
        "jacketing_member_report": jacketing_member_report_df,
        "failed_member_diagnosis": failed_member_diagnosis_df,
        "figures": figures,
    }
