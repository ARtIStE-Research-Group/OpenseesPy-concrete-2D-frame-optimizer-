import math
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt

from framekit.constants import (
    DEFAULT_GAMMA_C,
    DEFAULT_GAMMA_S,
    DEFAULT_ALPHA_CC,
    BAND_GREEN,
    BAND_YELLOW,
    BAND_ORANGE,
)
from framekit.data_models import EC2MaterialStrengths
from framekit.input_helpers import ask_choice, ask_float, ask_yes_no
from framekit.opensees_model import ops_ele_end_forces_local


# ============================================================
# GENERIC SECTION ACCESS HELPERS
# supports both old names (b, h, top_bars, ...) and newer ones
# ============================================================
def _get_attr(sec, *names, default=None):
    for name in names:
        if hasattr(sec, name):
            return getattr(sec, name)
    return default


def _sec_b(sec):
    return float(_get_attr(sec, "b", "b_mm", default=0.0))


def _sec_h(sec):
    return float(_get_attr(sec, "h", "h_mm", default=0.0))


def _sec_cover(sec):
    return float(_get_attr(sec, "cover", "cover_mm", default=0.0))


def _sec_top_bars(sec):
    return int(_get_attr(sec, "top_bars", "n_top", default=0))


def _sec_bottom_bars(sec):
    return int(_get_attr(sec, "bottom_bars", "n_bot", default=0))


def _sec_top_bar_d(sec):
    return float(_get_attr(sec, "top_bar_d", "phi_top_mm", default=0.0))


def _sec_bottom_bar_d(sec):
    return float(_get_attr(sec, "bottom_bar_d", "phi_bot_mm", default=0.0))


def _sec_role(sec):
    return str(_get_attr(sec, "section_role", default="beam") or "beam").lower()


def _sec_left_bars(sec):
    return int(_get_attr(sec, "left_bars", "n_left", default=0) or 0)


def _sec_right_bars(sec):
    return int(_get_attr(sec, "right_bars", "n_right", default=0) or 0)


def _sec_side_bar_d(sec):
    val = _get_attr(sec, "side_bar_d", "phi_side_mm", default=None)
    if val is None:
        vals = [v for v in [_sec_top_bar_d(sec), _sec_bottom_bar_d(sec)] if v and v > 0]
        return float(vals[0]) if vals else 0.0
    return float(val)


def _sec_is_column_like(sec, etype=None):
    if etype is not None and str(etype).lower() == "column":
        return True
    return _sec_role(sec) == "column"


def _jacket(sec):
    """(n_top, n_bot, n_left, n_right, phi, y_edge) of the jacket bars, or None."""
    t = float(_get_attr(sec, "jacket_t_mm", default=0.0) or 0.0)
    phi = float(_get_attr(sec, "jphi_mm", default=0.0) or 0.0)
    if t <= 0 or phi <= 0:
        return None
    nt = int(_get_attr(sec, "jn_top", default=0) or 0)
    nb = int(_get_attr(sec, "jn_bot", default=0) or 0)
    nl = int(_get_attr(sec, "jn_left", default=0) or 0)
    nr = int(_get_attr(sec, "jn_right", default=0) or 0)
    e = (float(_get_attr(sec, "jacket_cover_mm", default=30.0)) +
         float(_get_attr(sec, "jacket_stirrup_phi_mm", default=8.0)) + 0.5 * phi)
    return nt, nb, nl, nr, phi, e


def _tension_steel(sec, bottom=True):
    """Tension steel area [mm2] and effective depth d [mm] (existing + jacket bars)."""
    h = _sec_h(sec)
    cover = _sec_cover(sec)
    st = _sec_stirrup_phi(sec)
    n = _sec_bottom_bars(sec) if bottom else _sec_top_bars(sec)
    phi = _sec_bottom_bar_d(sec) if bottom else _sec_top_bar_d(sec)
    layers = []
    if n > 0 and phi > 0:
        layers.append((n * _bar_area_mm2(phi), h - (cover + st + 0.5 * phi)))
    j = _jacket(sec)
    if j is not None:
        nt, nb, nl, nr, jphi, e = j
        nj = nb if bottom else nt
        if nj > 0:
            layers.append((nj * _bar_area_mm2(jphi), h - e))
    As = sum(a for a, _ in layers)
    if As <= 0:
        return 0.0, h - cover - st
    d = sum(a * dd for a, dd in layers) / As
    return As, d


def rc_bar_areas(sec, etype=None):
    """Return longitudinal steel areas by position.

    For beams, side bars are zero. For columns, n_left/n_right are
    intermediate side-face bars excluding corner bars, as defined in
    section_library.py.
    """
    As_top = _sec_top_bars(sec) * _bar_area_mm2(_sec_top_bar_d(sec))
    As_bot = _sec_bottom_bars(sec) * _bar_area_mm2(_sec_bottom_bar_d(sec))
    if _sec_is_column_like(sec, etype):
        As_side = (_sec_left_bars(sec) + _sec_right_bars(sec)) * _bar_area_mm2(_sec_side_bar_d(sec))
    else:
        As_side = 0.0
    j = _jacket(sec)
    if j is not None:
        nt, nb, nl, nr, jphi, e = j
        As_top += nt * _bar_area_mm2(jphi)
        As_bot += nb * _bar_area_mm2(jphi)
        As_side += (nl + nr) * _bar_area_mm2(jphi)
    return {"As_top_mm2": As_top, "As_bottom_mm2": As_bot, "As_side_mm2": As_side, "As_total_mm2": As_top + As_bot + As_side}


def rc_reinforcement_ratio(sec, etype=None):
    Ac = max(_sec_b(sec) * _sec_h(sec), 1e-9)
    return rc_bar_areas(sec, etype=etype)["As_total_mm2"] / Ac


def _sec_stirrup_phi(sec):
    return float(_get_attr(sec, "stirrup_phi_mm", "stirrup_phi", default=8.0))


def _sec_stirrup_spacing(sec):
    return float(_get_attr(sec, "stirrup_spacing_mm", "stirrup_spacing", default=150.0))


def _sec_stirrup_spacing_critical(sec):
    return float(_get_attr(sec, "stirrup_spacing_critical_mm", "stirrup_spacing_critical", default=_sec_stirrup_spacing(sec)))


def _sec_stirrup_legs(sec):
    return int(_get_attr(sec, "stirrup_legs", default=2))


def _bar_area_mm2(phi_mm):
    return math.pi * float(phi_mm) ** 2 / 4.0


# ============================================================
# EC2 MATERIAL INPUT
# ============================================================
def ask_ec2_materials_default_or_custom():
    print("\n--- EC2 material strengths for checks ---")
    print("1) Use defaults: C25/30 and B500")
    print("2) Enter fck and fyk")
    opt = ask_choice("Select option (1/2): ", ["1", "2"])

    if opt == "1":
        fck = 25.0
        fyk = 500.0
    else:
        fck = ask_float("Concrete fck [MPa]: ", min_val=12.0, allow_zero=False)
        fyk = ask_float("Steel fyk [MPa]: ", min_val=200.0, allow_zero=False)

    gamma_c = DEFAULT_GAMMA_C
    gamma_s = DEFAULT_GAMMA_S
    alpha_cc = DEFAULT_ALPHA_CC

    if ask_yes_no(f"Change gamma_c? (default {DEFAULT_GAMMA_C}) (y/n): "):
        gamma_c = ask_float("gamma_c [-]: ", min_val=1.0, allow_zero=False)

    if ask_yes_no(f"Change gamma_s? (default {DEFAULT_GAMMA_S}) (y/n): "):
        gamma_s = ask_float("gamma_s [-]: ", min_val=1.0, allow_zero=False)

    if ask_yes_no(f"Change alpha_cc? (default {DEFAULT_ALPHA_CC}) (y/n): "):
        alpha_cc = ask_float("alpha_cc [-]: ", min_val=0.5, max_val=1.0, allow_zero=False)

    return EC2MaterialStrengths(
        fck=fck,
        fyk=fyk,
        gamma_c=gamma_c,
        gamma_s=gamma_s,
        alpha_cc=alpha_cc
    )


# ============================================================
# BASIC RC HELPERS
# ============================================================
def rc_As_total(sec):
    """Total longitudinal steel area in mm2.

    Backward-compatible for old beam sections, but column-aware when the
    section has section_role='column' and side-bar fields n_left/n_right.
    """
    return rc_bar_areas(sec, etype=None)["As_total_mm2"]


def ec2_fcd(mat):
    return mat.alpha_cc * mat.fck / mat.gamma_c


def ec2_fyd(mat):
    return mat.fyk / mat.gamma_s


def ec2_fywd(mat):
    return mat.fyk / mat.gamma_s


def compute_transverse_reinf_report(sec, mat=None):
    stirrup_phi = _sec_stirrup_phi(sec)
    stirrup_spacing = _sec_stirrup_spacing(sec)
    stirrup_spacing_critical = _sec_stirrup_spacing_critical(sec)
    stirrup_legs = _sec_stirrup_legs(sec)

    Asw_leg = _bar_area_mm2(stirrup_phi)
    Asw_total = stirrup_legs * Asw_leg
    Asw_per_s_mm2_per_mm = Asw_total / max(stirrup_spacing, 1e-9)
    Asw_per_s_mm2_per_m = 1000.0 * Asw_per_s_mm2_per_mm

    fywd = ec2_fywd(mat) if mat is not None else 500.0 / 1.15

    return {
        "stirrup_phi_mm": stirrup_phi,
        "stirrup_spacing_mm": stirrup_spacing,
        "stirrup_spacing_critical_mm": stirrup_spacing_critical,
        "stirrup_legs": stirrup_legs,
        "Asw_leg_mm2": Asw_leg,
        "Asw_total_mm2": Asw_total,
        "Asw_per_s_mm2_per_m": Asw_per_s_mm2_per_m,
        "fywd_MPa": fywd,
    }


def ec2_MRd_rectangular(sec, mat, sagging=True):
    b = _sec_b(sec)
    h = _sec_h(sec)
    cover = _sec_cover(sec)

    fcd = ec2_fcd(mat)
    fyd = ec2_fyd(mat)

    As, d = _tension_steel(sec, bottom=bool(sagging))

    if As <= 0 or d <= 0 or b <= 0 or h <= 0:
        return 0.0

    T = As * fyd
    x = T / max(0.8 * b * fcd, 1e-9)
    x_lim = 0.45 * d
    if x > x_lim:
        # concrete block governs: limit compression force to 0.8*b*x_lim*fcd
        x = x_lim
        T = 0.8 * b * x * fcd

    z = d - 0.4 * x
    z = max(min(z, 0.95 * d), 0.2 * d)

    MRd_Nmm = T * z
    return MRd_Nmm / 1e6  # kNm


def ec2_NRd_compression(sec, mat):
    b = _sec_b(sec)
    h = _sec_h(sec)
    Ac = b * h
    As = rc_As_total(sec)

    fcd = ec2_fcd(mat)
    fyd = ec2_fyd(mat)

    # EC2 6.1(5): pure compression strain eps_c2 = 2 permil -> sigma_s <= 0.002*Es
    NRd_N = fcd * Ac + min(fyd, 0.002 * 200000.0) * As
    return NRd_N / 1000.0  # kN


def ec2_VRd_concrete(sec, mat, NEd_kN=0.0):
    fck = mat.fck
    fcd = ec2_fcd(mat)

    bw = _sec_b(sec)
    h = _sec_h(sec)
    cover = _sec_cover(sec)
    As_l, d = _tension_steel(sec, bottom=True)
    d = max(d, 0.6 * h)

    rho_l = min(As_l / max(bw * d, 1e-9), 0.02)

    k = min(1.0 + math.sqrt(200.0 / max(d, 1e-9)), 2.0)

    Ac = bw * h
    NEd_N = NEd_kN * 1000.0
    sigma_cp = min(NEd_N / max(Ac, 1e-9), 0.2 * fcd)

    CRd_c = 0.18 / mat.gamma_c
    k1 = 0.15

    VRd_c_expr = (CRd_c * k * (100.0 * rho_l * fck) ** (1.0 / 3.0) + k1 * sigma_cp) * bw * d
    v_min = 0.035 * k ** 1.5 * math.sqrt(fck)
    VRd_c_min = (v_min + k1 * sigma_cp) * bw * d
    VRd_c = max(VRd_c_expr, VRd_c_min, 0.0)

    nu1 = 0.6 * (1.0 - fck / 250.0)
    z = 0.9 * d
    VRd_max = nu1 * fcd * bw * z

    VRd = min(VRd_c, VRd_max)
    return VRd / 1000.0  # kN


def ec2_VRd_stirrups(sec, mat):
    rep = compute_transverse_reinf_report(sec, mat)

    bw = _sec_b(sec)
    h = _sec_h(sec)
    cover = _sec_cover(sec)
    As_l, d = _tension_steel(sec, bottom=True)
    d = max(d, 0.6 * h)
    z = 0.9 * d

    Asw_per_s = rep["Asw_total_mm2"] / max(rep["stirrup_spacing_mm"], 1e-9)  # mm2/mm
    fywd = rep["fywd_MPa"]

    # simplified EC2 shear reinforcement contribution with theta = 45deg -> cot(theta)=1
    VRds_N = Asw_per_s * z * fywd
    nu1 = 0.6 * (1.0 - mat.fck / 250.0)
    VRd_max_N = 1.0 * bw * z * nu1 * ec2_fcd(mat) / (1.0 + 1.0)
    VRds_kN = min(VRds_N, VRd_max_N) / 1000.0

    return {
        **rep,
        "d_mm": d,
        "z_mm": z,
        "VRds_kN": VRds_kN
    }


def ec2_VRd(sec, mat, NEd_kN=0.0):
    VRdc = ec2_VRd_concrete(sec, mat, NEd_kN=NEd_kN)
    stir = ec2_VRd_stirrups(sec, mat)
    VRds = stir["VRds_kN"]

    # for this modular simplified check we use the larger resisting contribution
    # this keeps compatibility with previous behavior while including stirrup effect
    return max(VRdc, VRds)


# ============================================================
# COLUMN N-M INTERACTION
# ============================================================
def _column_bar_layers_from_top(sec):
    """Return [(area_mm2, y_from_top_mm), ...] for column longitudinal bars."""
    h = _sec_h(sec)
    cover = _sec_cover(sec)
    stirrup_phi = _sec_stirrup_phi(sec)

    n_top = _sec_top_bars(sec)
    n_bot = _sec_bottom_bars(sec)
    phi_top = _sec_top_bar_d(sec)
    phi_bot = _sec_bottom_bar_d(sec)
    n_left = _sec_left_bars(sec)
    n_right = _sec_right_bars(sec)
    phi_side = _sec_side_bar_d(sec)

    y_top = cover + stirrup_phi + 0.5 * phi_top if n_top > 0 else cover
    y_bot = h - (cover + stirrup_phi + 0.5 * phi_bot) if n_bot > 0 else h - cover

    layers = []
    if n_top > 0 and phi_top > 0:
        layers.append((n_top * _bar_area_mm2(phi_top), y_top))
    if n_bot > 0 and phi_bot > 0:
        layers.append((n_bot * _bar_area_mm2(phi_bot), y_bot))

    # Side bars are intermediate bars along left and right faces. For uniaxial
    # bending they are represented by their vertical positions and tributary
    # steel areas. This lets side bars contribute to both axial and N-M capacity.
    n_side_total = n_left + n_right
    if n_side_total > 0 and phi_side > 0 and y_bot > y_top:
        # Use one layer for each elevation; area is sum of left/right bars at that elevation.
        n_levels = max(n_left, n_right)
        for i in range(n_levels):
            y = y_top + (i + 1) * (y_bot - y_top) / (n_levels + 1)
            n_at_level = (1 if i < n_left else 0) + (1 if i < n_right else 0)
            if n_at_level:
                layers.append((n_at_level * _bar_area_mm2(phi_side), y))
    j = _jacket(sec)
    if j is not None:
        nt, nb, nl, nr, jphi, e = j
        a = _bar_area_mm2(jphi)
        if nt:
            layers.append((nt * a, e))
        if nb:
            layers.append((nb * a, h - e))
        for i in range(max(nl, nr)):
            y = e + (i + 1) * (h - 2 * e) / (max(nl, nr) + 1)
            k = (1 if i < nl else 0) + (1 if i < nr else 0)
            layers.append((k * a, y))
    return layers


def ec2_MRd_column_all_bars(sec, mat):
    """Simplified uniaxial flexural resistance for columns using all perimeter bars."""
    fcd = ec2_fcd(mat)
    fyd = ec2_fyd(mat)
    b = _sec_b(sec)
    h = _sec_h(sec)
    if b <= 0 or h <= 0:
        return 0.0

    layers = _column_bar_layers_from_top(sec)
    if not layers:
        return ec2_MRd_rectangular(sec, mat, sagging=True)

    eps_cu = 0.0035
    Es = 200000.0
    lambda_ = 0.8
    eta_cc = 1.0

    best = 0.0
    # Sweep neutral-axis depths; include shallow and deep compression blocks.
    xs = list(np.linspace(0.05 * h, 2.5 * h, 160))
    for x in xs:
        a = min(lambda_ * x, h)
        Fc = eta_cc * fcd * b * a
        M = Fc * (h / 2.0 - a / 2.0)
        for As, y in layers:
            eps_s = eps_cu * (x - y) / max(x, 1e-9)
            sig_s = max(min(Es * eps_s, fyd), -fyd)
            Fs = As * sig_s
            M += Fs * (h / 2.0 - y)
        best = max(best, abs(M) / 1e6)
    return best


def ec2_NM_interaction_column(sec, mat, NEd_kN, MEd_kNm):
    """Column N-M check with all-around longitudinal reinforcement.

    The older implementation used only top and bottom bars. This version also
    includes n_left/n_right side-face bars in axial resistance and in a simplified
    strain-compatible uniaxial N-M interaction curve. It remains intentionally
    lightweight for the framework, but the side bars now affect eta_NM.
    """
    fcd = ec2_fcd(mat)
    fyd = ec2_fyd(mat)
    b = _sec_b(sec)
    h = _sec_h(sec)

    NEd = float(NEd_kN)
    MEd = abs(float(MEd_kNm))

    layers = _column_bar_layers_from_top(sec)
    if b <= 0 or h <= 0 or not layers:
        MRd0_kNm = ec2_MRd_rectangular(sec, mat, sagging=True)
        NRd0_kN = ec2_NRd_compression(sec, mat)
        eta = max(MEd / max(MRd0_kNm, 1e-9), NEd / max(NRd0_kN, 1e-9) if NEd > 0 else 0.0)
        return {
            "NRd0_kN": NRd0_kN,
            "MRd0_kNm": MRd0_kNm,
            "Nbal_kN": np.nan,
            "MRd_bal_kNm": MRd0_kNm,
            "eta_NM": float(eta),
            "ok": float(eta) <= 1.0,
        }

    NRd0_kN = ec2_NRd_compression(sec, mat)
    MRd0_kNm = ec2_MRd_column_all_bars(sec, mat)

    eps_cu = 0.0035
    Es = 200000.0
    lambda_ = 0.8
    eta_cc = 1.0

    points = [(NRd0_kN, 0.0)]
    # Build a positive-compression interaction curve. N is compression-positive.
    for x in np.linspace(0.03 * h, 3.0 * h, 260):
        a = min(lambda_ * x, h)
        Fc = eta_cc * fcd * b * a
        N = Fc
        M = Fc * (h / 2.0 - a / 2.0)
        for As, y in layers:
            eps_s = eps_cu * (x - y) / max(x, 1e-9)
            sig_s = max(min(Es * eps_s, fyd), -fyd)
            Fs = As * sig_s
            N += Fs
            M += Fs * (h / 2.0 - y)
        N_kN = N / 1000.0
        M_kNm = abs(M) / 1e6
        if np.isfinite(N_kN) and np.isfinite(M_kNm) and N_kN >= -0.30 * NRd0_kN:
            points.append((N_kN, M_kNm))

    # Remove pathological/duplicate points and order from high N to low N.
    clean = []
    seen = set()
    for N, M in points:
        key = (round(float(N), 3), round(float(M), 3))
        if key not in seen and M >= 0.0:
            clean.append((float(N), float(M)))
            seen.add(key)
    clean = sorted(clean, key=lambda p: p[0], reverse=True)

    def _segment_eta(Nx, Mx, N1, M1, N2, M2):
        dN = N2 - N1
        dM = M2 - M1
        denom = Nx * dM - Mx * dN
        if abs(denom) < 1e-12:
            return None
        t = (N1 * dM - M1 * dN) / denom
        if abs(dN) > 1e-12:
            s = (t * Nx - N1) / dN
        elif abs(dM) > 1e-12:
            s = (t * Mx - M1) / dM
        else:
            return None
        if t < 1e-12 or s < -1e-9 or s > 1.0 + 1e-9:
            return None
        return 1.0 / t

    eta_candidates = []
    for (N1, M1), (N2, M2) in zip(clean[:-1], clean[1:]):
        val = _segment_eta(NEd, MEd, N1, M1, N2, M2)
        if val is not None and np.isfinite(val) and val >= 0.0:
            eta_candidates.append(val)

    if eta_candidates:
        eta = min(eta_candidates)
    else:
        # Robust fallback: interaction by separate axial/flexural ratios.
        eta = MEd / max(MRd0_kNm, 1e-9)
        if NEd > 0.0:
            eta = max(eta, NEd / max(NRd0_kN, 1e-9))

    # Balanced point only for reporting: take maximum moment point on curve.
    if clean:
        Nbal_kN, MRd_bal_kNm = max(clean, key=lambda p: p[1])
    else:
        Nbal_kN, MRd_bal_kNm = np.nan, MRd0_kNm

    return {
        "NRd0_kN": NRd0_kN,
        "MRd0_kNm": MRd0_kNm,
        "Nbal_kN": Nbal_kN,
        "MRd_bal_kNm": MRd_bal_kNm,
        "eta_NM": float(eta),
        "ok": float(eta) <= 1.0,
    }


# ============================================================
# DEMANDS
# ============================================================
def element_demands_from_local_force(localF):
    N1, V1, M1, N2, V2, M2 = localF
    NEd = max(abs(N1), abs(N2)) / 1000.0
    VEd = max(abs(V1), abs(V2)) / 1000.0
    MEd = max(abs(M1), abs(M2)) / 1000.0
    # OpenSees localForce: compression at end i -> N1 > 0, at end j -> N2 < 0.
    # Most compressive axial force, compression positive (negative = tension).
    NEd_signed = max(N1, -N2) / 1000.0
    return NEd, VEd, MEd, NEd_signed


def ec2_check_element(sec, etype, mat, NEd_kN, VEd_kN, MEd_kNm, NEd_signed_kN=0.0,
                      MEd_sag_kNm=None, MEd_hog_kNm=None):
    stirrup_info = compute_transverse_reinf_report(sec, mat)
    steel_info = rc_bar_areas(sec, etype=etype)
    steel_info["rho_long_pct"] = 100.0 * rc_reinforcement_ratio(sec, etype=etype)
    VRd = ec2_VRd(sec, mat, NEd_kN=NEd_signed_kN if etype == "column" else 0.0)
    VRdc = ec2_VRd_concrete(sec, mat, NEd_kN=NEd_signed_kN if etype == "column" else 0.0)
    VRds = ec2_VRd_stirrups(sec, mat)["VRds_kN"]

    etaV = VEd_kN / max(VRd, 1e-9)

    if etype == "beam":
        MRd_sag = ec2_MRd_rectangular(sec, mat, sagging=True)
        MRd_hog = ec2_MRd_rectangular(sec, mat, sagging=False)
        MRd = max(MRd_sag, MRd_hog)
        NRd = ec2_NRd_compression(sec, mat)

        if MEd_sag_kNm is not None and MEd_hog_kNm is not None:
            etaM = max(MEd_sag_kNm / max(MRd_sag, 1e-9),
                       MEd_hog_kNm / max(MRd_hog, 1e-9))
        else:
            # sign unknown: check against the weaker face
            etaM = MEd_kNm / max(min(MRd_sag, MRd_hog), 1e-9)
        etaN = NEd_kN / max(NRd, 1e-9)
        eta_NM = etaM
        eta = max(etaV, eta_NM)

        return {
            "NRd_kN": NRd,
            "VRd_kN": VRd,
            "VRdc_kN": VRdc,
            "VRds_kN": VRds,
            "MRd_kNm": MRd,
            "etaN": etaN,
            "etaV": etaV,
            "etaM": etaM,
            "eta_NM": eta_NM,
            "eta": eta,
            "NM_detail": None,
            **steel_info,
            **stirrup_info,
        }

    nm = ec2_NM_interaction_column(sec, mat, NEd_kN=NEd_signed_kN, MEd_kNm=MEd_kNm)
    eta_NM = nm["eta_NM"]
    eta = max(etaV, eta_NM)

    return {
        "NRd_kN": nm["NRd0_kN"],
        "VRd_kN": VRd,
        "VRdc_kN": VRdc,
        "VRds_kN": VRds,
        "MRd_kNm": nm["MRd_bal_kNm"],
        "etaN": NEd_kN / max(nm["NRd0_kN"], 1e-9),
        "etaV": etaV,
        "etaM": MEd_kNm / max(nm["MRd0_kNm"], 1e-9),
        "eta_NM": eta_NM,
        "eta": eta,
        "NM_detail": nm,
        **steel_info,
        **stirrup_info,
    }


# ============================================================
# RUN CHECKS FOR ONE COMBINATION
# ============================================================
def run_combo_and_check_all(
    nodes,
    elements,
    loads_combo,
    Ec_Pa,
    Es_Pa,
    column_sec,
    beam_sec,
    ec2mat,
    ops_build_model,
    ops_apply_combo_loads,
    ops_run_static,
    geom=None,
    use_rigid_diaphragm=False,
    walls=None,
    wall_tw_mm=200.0,
    wall_E_Pa=30e9,
    element_sections=None,
):
    ops_build_model(
        nodes,
        elements,
        Ec_Pa,
        Es_Pa,
        column_sec,
        beam_sec,
        geom=geom,
        use_rigid_diaphragm=use_rigid_diaphragm,
        walls=walls,
        wall_tw_mm=wall_tw_mm,
        wall_E_Pa=wall_E_Pa,
        element_sections=element_sections
    )

    ops_apply_combo_loads(nodes, elements, loads_combo)
    ok = ops_run_static()

    if ok != 0:
        return None, ok

    checks = {}
    for eid, (_, _, etype) in elements.items():
        if element_sections is not None and int(eid) in element_sections:
            sec = element_sections[int(eid)]
        else:
            sec = column_sec if etype == "column" else beam_sec
        lf = ops_ele_end_forces_local(eid)
        NEd, VEd, MEd, NEd_signed = element_demands_from_local_force(lf)

        MEd_sag = MEd_hog = None
        if etype == "beam":
            # span moments (sagging +) incl. distributed loads, not only end moments
            from framekit.diagrams import _nvm_arrays_for_element
            res = _nvm_arrays_for_element(eid, elements, loads_combo, nodes, npts=41)
            if res is not None:
                M_kNm = np.asarray(res[4], dtype=float) / 1000.0
                MEd_sag = max(float(M_kNm.max()), 0.0)
                MEd_hog = max(float(-M_kNm.min()), 0.0)
                MEd = max(MEd, MEd_sag, MEd_hog)

        checks[eid] = {
            "etype": etype,
            "NEd_kN": NEd,
            "VEd_kN": VEd,
            "MEd_kNm": MEd,
            "NEd_signed_kN": NEd_signed,
            **ec2_check_element(sec, etype, ec2mat, NEd, VEd, MEd, NEd_signed_kN=NEd_signed,
                                MEd_sag_kNm=MEd_sag, MEd_hog_kNm=MEd_hog)
        }

    return checks, 0


# ============================================================
# WORST TABLES
# ============================================================
def utilization_to_color(eta):
    e = float(eta)
    if e <= BAND_GREEN:
        return (0.75, 0.95, 0.75, 1.0)
    if e <= BAND_YELLOW:
        return (1.00, 1.00, 0.70, 1.0)
    if e <= BAND_ORANGE:
        return (1.00, 0.85, 0.60, 1.0)
    return (1.00, 0.70, 0.70, 1.0)


def build_worst_table(checks_by_combo):
    worst = {}
    for cname, checks in checks_by_combo.items():
        for eid, r in checks.items():
            if eid not in worst or r["eta"] > worst[eid]["eta"]:
                worst[eid] = {**r, "combo": cname}

    rows = sorted([(eid, data) for eid, data in worst.items()], key=lambda x: x[1]["eta"], reverse=True)
    return rows


def worst_table_to_dataframe(checks_by_combo):
    rows = []
    worst = build_worst_table(checks_by_combo)

    for eid, d in worst:
        rows.append({
            "Element": eid,
            "ElementID": eid,
            "Type": d["etype"],
            "Combination": d["combo"],
            "eta": d["eta"],
            "eta_NM": d["eta_NM"],
            "etaV": d["etaV"],
            "etaM": d["etaM"],
            "etaN": d["etaN"],
            "NEd_kN": d["NEd_kN"],
            "VEd_kN": d["VEd_kN"],
            "MEd_kNm": d["MEd_kNm"],
            "NRd_kN": d["NRd_kN"],
            "VRd_kN": d["VRd_kN"],
            "VRdc_kN": d.get("VRdc_kN", None),
            "VRds_kN": d.get("VRds_kN", None),
            "MRd_kNm": d["MRd_kNm"],
            "As_top_mm2": d.get("As_top_mm2", None),
            "As_bottom_mm2": d.get("As_bottom_mm2", None),
            "As_side_mm2": d.get("As_side_mm2", None),
            "As_total_mm2": d.get("As_total_mm2", None),
            "rho_long_pct": d.get("rho_long_pct", None),
            "StaticPass": "OK" if d["eta"] <= 1.0 else "FAIL",
            "stirrup_phi_mm": d.get("stirrup_phi_mm", None),
            "stirrup_spacing_mm": d.get("stirrup_spacing_mm", None),
            "stirrup_spacing_critical_mm": d.get("stirrup_spacing_critical_mm", None),
            "stirrup_legs": d.get("stirrup_legs", None),
            "Asw_leg_mm2": d.get("Asw_leg_mm2", None),
            "Asw_total_mm2": d.get("Asw_total_mm2", None),
            "Asw_per_s_mm2_per_m": d.get("Asw_per_s_mm2_per_m", None),
            "fywd_MPa": d.get("fywd_MPa", None),
        })

    return pd.DataFrame(rows)


def add_stirrup_columns_to_worst_table(df_worst, beam_sec=None, column_sec=None, elements=None):
    # kept for compatibility with your main.py
    # if the columns already exist, just return the dataframe
    if df_worst is None or df_worst.empty:
        return df_worst

    needed = {
        "stirrup_phi_mm",
        "stirrup_spacing_mm",
        "stirrup_legs",
        "Asw_per_s_mm2_per_m",
    }
    if needed.issubset(set(df_worst.columns)):
        return df_worst

    df = df_worst.copy()
    for c in [
        "stirrup_phi_mm",
        "stirrup_spacing_mm",
        "stirrup_spacing_critical_mm",
        "stirrup_legs",
        "Asw_leg_mm2",
        "Asw_total_mm2",
        "Asw_per_s_mm2_per_m",
    ]:
        if c not in df.columns:
            df[c] = None

    if elements is None or beam_sec is None or column_sec is None:
        return df

    ele_col = "Element" if "Element" in df.columns else ("ElementID" if "ElementID" in df.columns else None)
    if ele_col is None:
        return df

    for i, row in df.iterrows():
        eid = int(row[ele_col])
        if eid not in elements:
            continue

        _, _, etype = elements[eid]
        sec = beam_sec if etype == "beam" else column_sec
        rep = compute_transverse_reinf_report(sec)

        df.at[i, "stirrup_phi_mm"] = rep["stirrup_phi_mm"]
        df.at[i, "stirrup_spacing_mm"] = rep["stirrup_spacing_mm"]
        df.at[i, "stirrup_spacing_critical_mm"] = rep["stirrup_spacing_critical_mm"]
        df.at[i, "stirrup_legs"] = rep["stirrup_legs"]
        df.at[i, "Asw_leg_mm2"] = rep["Asw_leg_mm2"]
        df.at[i, "Asw_total_mm2"] = rep["Asw_total_mm2"]
        df.at[i, "Asw_per_s_mm2_per_m"] = rep["Asw_per_s_mm2_per_m"]

    return df


# ============================================================
# COLORED TABLE PLOT
# ============================================================
def plot_verification_table(checks_by_combo, title="EC2 Verification (worst utilization per member)", max_rows=40):
    from framekit.plot_style import plot_table
    rows = build_worst_table(checks_by_combo)[:max_rows]
    headers = ["Member", "Type", "Governing combo", "η", "η N-M", "η V", "NEd [kN]", "VEd [kN]", "MEd [kNm]"]
    cell_text = []
    for eid, d in rows:
        cell_text.append([
            str(eid), d["etype"], d["combo"],
            f"{d['eta']:.2f}", f"{d.get('eta_NM', 0.0):.2f}", f"{d['etaV']:.2f}",
            f"{d['NEd_kN']:.1f}", f"{d['VEd_kN']:.1f}", f"{d['MEd_kNm']:.1f}",
        ])
    plot_table(cell_text, headers, title, eta_cols=("η", "η N-M", "η V"), max_rows=max_rows,
               col_widths=[0.07, 0.08, 0.2, 0.08, 0.08, 0.08, 0.11, 0.11, 0.11])


# ============================================================
# COLORED STRUCTURE UTILIZATION PLOT
# ============================================================
def plot_utilization_on_structure(nodes, elements, checks_by_combo, title="Utilization map (worst eta)",
                                  geom=None, walls=None, info=None):
    from framekit.plot_style import plot_utilization_frame
    eta_map = {eid: d["eta"] for eid, d in build_worst_table(checks_by_combo)}
    plot_utilization_frame(nodes, elements, eta_map, title=title, geom=geom, walls=walls,
                           info=info if info is not None else "envelope of " + ", ".join(sorted(checks_by_combo)))


# ============================================================
# COLORED EXCEL EXPORT
# ============================================================
def export_ec2_worst_table_colored(df, filename="ec2_worst_utilization.xlsx"):
    if df.empty:
        print("❌ No EC2 table to export.")
        return

    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment

        wb = Workbook()
        ws = wb.active
        ws.title = "EC2_Worst"

        headers = list(df.columns)
        ws.append(headers)

        header_fill = PatternFill(fill_type="solid", start_color="D9D9D9", end_color="D9D9D9")
        bold_font = Font(bold=True)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = bold_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        fill_green = PatternFill(fill_type="solid", start_color="C6EFCE", end_color="C6EFCE")
        fill_yellow = PatternFill(fill_type="solid", start_color="FFF2CC", end_color="FFF2CC")
        fill_orange = PatternFill(fill_type="solid", start_color="FCE4D6", end_color="FCE4D6")
        fill_red = PatternFill(fill_type="solid", start_color="F4CCCC", end_color="F4CCCC")

        eta_idx = headers.index("eta")
        static_pass_idx = headers.index("StaticPass") if "StaticPass" in headers else None

        for _, row in df.iterrows():
            ws.append(list(row))

        for i in range(2, ws.max_row + 1):
            eta = ws.cell(row=i, column=eta_idx + 1).value

            if eta is None:
                continue

            if eta <= BAND_GREEN:
                fill = fill_green
            elif eta <= BAND_YELLOW:
                fill = fill_yellow
            elif eta <= BAND_ORANGE:
                fill = fill_orange
            else:
                fill = fill_red

            for j in range(1, ws.max_column + 1):
                cell = ws.cell(row=i, column=j)
                cell.fill = fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            if static_pass_idx is not None:
                pass_cell = ws.cell(row=i, column=static_pass_idx + 1)
                if str(pass_cell.value).upper() == "OK":
                    pass_cell.fill = fill_green
                else:
                    pass_cell.fill = fill_red

        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max_len + 2

        wb.save(filename)
        print(f"✔ Colored Excel exported: {filename}")

    except Exception as e:
        print(f"❌ Excel export failed: {e}")